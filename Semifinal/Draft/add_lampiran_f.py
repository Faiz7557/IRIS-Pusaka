# -*- coding: utf-8 -*-
"""Tambahkan Lampiran F (13 slide, bagian lanskap satu kolom) ke akhir naskah."""
import copy, sys, os
import docx
from docx.shared import Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC, DST, PNG_DIR = sys.argv[1], sys.argv[2], sys.argv[3]
d = docx.Document(SRC)
body = d.element.body
final_sect = body.find(qn("w:sectPr"))

# 1) tutup bagian potret 2-kolom lama pada paragraf kosong bersectPr
old = copy.deepcopy(final_sect)
p_end = OxmlElement("w:p")
ppr = OxmlElement("w:pPr"); ppr.append(old); p_end.append(ppr)
final_sect.addprevious(p_end)

# 2) ubah sectPr akhir menjadi lanskap 1 kolom, mulai halaman baru
for t in final_sect.findall(qn("w:type")):
    final_sect.remove(t)
typ = OxmlElement("w:type"); typ.set(qn("w:val"), "nextPage")
# w:type harus setelah header/footer reference
last_ref = [c for c in final_sect if c.tag in (qn("w:headerReference"), qn("w:footerReference"))]
(last_ref[-1] if last_ref else final_sect[0]).addnext(typ) if last_ref else final_sect.insert(0, typ)
pg = final_sect.find(qn("w:pgSz")); pg.set(qn("w:w"), "15840"); pg.set(qn("w:h"), "12240"); pg.set(qn("w:orient"), "landscape")
cols = final_sect.find(qn("w:cols"))
for c in list(cols):
    cols.remove(c)
for a in list(cols.attrib):
    del cols.attrib[a]
cols.set(qn("w:space"), "720")

# 3) isi
def par(text_runs, style=None, jc=None, before=None, after=None, keep_next=False, page_break=False, size=None, italic=False):
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    if style:
        ps = OxmlElement("w:pStyle"); ps.set(qn("w:val"), style); ppr.append(ps)
    if keep_next:
        ppr.append(OxmlElement("w:keepNext"))
    if page_break:
        ppr.append(OxmlElement("w:pageBreakBefore"))
    if before is not None or after is not None:
        sp = OxmlElement("w:spacing")
        if before is not None: sp.set(qn("w:before"), str(before))
        if after is not None: sp.set(qn("w:after"), str(after))
        ppr.append(sp)
    if style == "Heading1":
        ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "0"); ind.set(qn("w:firstLine"), "0"); ppr.append(ind)
    if jc:
        j = OxmlElement("w:jc"); j.set(qn("w:val"), jc); ppr.append(j)
    p.append(ppr)
    for txt, b in text_runs:
        r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
        if b: rpr.append(OxmlElement("w:b")); rpr.append(OxmlElement("w:bCs"))
        if italic: rpr.append(OxmlElement("w:i")); rpr.append(OxmlElement("w:iCs"))
        if size:
            s = OxmlElement("w:sz"); s.set(qn("w:val"), str(size)); rpr.append(s)
            s = OxmlElement("w:szCs"); s.set(qn("w:val"), str(size)); rpr.append(s)
        lg = OxmlElement("w:lang"); lg.set(qn("w:val"), "id-ID"); rpr.append(lg)
        r.append(rpr)
        t = OxmlElement("w:t"); t.text = txt; t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r.append(t); p.append(r)
    return p

SL = [
 ("Profil dan nilai jual SIMAK", "Ringkasan satu halaman: tiga bukti masalah, tiga langkah solusi, dan empat nilai jual. Status bukti dinyatakan apa adanya (terukur, teruji awal, belum terukur).",
  "Ringkasan tiga kolom: bukti masalah (1.727 mahasiswa tunanetra, 0 dari 28 PDF lulus PDF/UA-1, 2,6% gambar beralt), solusi SIMAK dalam tiga langkah, dan empat nilai jual, dengan bilah status bukti."),
 ("Pohon masalah dari tujuan hingga akar masalah dan buktinya", "Tiga kelompok masalah diturunkan menjadi enam akar masalah, masing-masing dengan bukti dari audit awal atau literatur; pohon ini mendasari tiga rumusan masalah pada Bab I.",
  "Diagram pohon: pertanyaan tujuan di kiri, tiga masalah utama beranak enam submasalah, akar masalah, dan bukti pada kolom kanan."),
 ("Value Proposition Canvas: kesesuaian solusi dengan pengguna", "Segmen utama adalah staf ULD; setiap pereda masalah pada peta nilai dipasangkan dengan masalah yang teramati pada profil pelanggan.",
  "Kanvas proposisi nilai: peta nilai SIMAK di kiri (produk, pereda masalah, pencipta manfaat) dan profil pelanggan di kanan (tugas, masalah, manfaat)."),
 ("Skenario ilustratif sebelum dan sesudah SIMAK", "Skenario disusun dari temuan audit awal; contoh deskripsi grafik bersifat ilustratif dan bukan keluaran sistem yang telah diuji.",
  "Perbandingan dua kolom pengalaman mahasiswa tunanetra sebelum dan sesudah SIMAK, dengan alur unggah, perakitan, dan gerbang verifikasi di bawahnya."),
 ("Analisis kompetitor: matriks pembanding fitur", "Matriks memperlihatkan letak pembeda SIMAK; sel diisi hanya dari yang dilaporkan sumber atau diamati penulis, dan “n/r” berarti tidak dilaporkan.",
  "Tabel pembanding tujuh solusi terhadap tujuh kriteria fitur, dengan kolom SIMAK disorot."),
 ("Segmentasi, target, dan posisi (STP)", "Segmen dan target dipilih dari kondisi yang terbukti pada audit awal; posisi pada peta merupakan penilaian kualitatif penulis.",
  "Empat dimensi segmentasi, empat prioritas target, peta posisi otomasi terhadap verifikasi, dan pernyataan posisi SIMAK."),
 ("Model bisnis dan keberlanjutan (Business Model Canvas ringkas)", "Pada tahap penelitian, sumber pendapatan dan saluran bersifat hipotesis; tidak ada angka biaya atau pendapatan yang diklaim.",
  "Sembilan blok kanvas model bisnis: mitra, aktivitas, sumber daya, proposisi nilai, hubungan, saluran, segmen, struktur biaya, dan sumber pendapatan."),
 ("Analisis SWOT", "Kelemahan dinyatakan apa adanya: penelitian berada pada tahap kelayakan dan rancangan evaluasi.",
  "Empat kuadran SWOT masing-masing berisi enam butir: kekuatan, kelemahan, peluang, dan ancaman."),
 ("Strategi TOWS", "Strategi diturunkan dari SWOT dan dijadwalkan pada peta jalan (Gambar F.13).",
  "Matriks dua kali dua strategi SO, ST, WO, dan WT, masing-masing tiga tindakan."),
 ("Lima kekuatan Porter", "Skor adalah penilaian awal tim, bukan hasil survei pasar; tekanan tertinggi berasal dari daya tawar pembeli dan persaingan antarpemain.",
  "Diagram radar lima kekuatan Porter dengan skor tiga, tiga, empat, empat, dan tiga dari lima, beserta penjelasan tiap kekuatan."),
 ("Analisis PESTEL", "Kondisi eksternal mendukung permintaan (kewajiban hukum, jumlah mahasiswa), sedangkan ketidakandalan model menjadikan verifikasi sebagai pembeda.",
  "Tabel enam faktor politik, ekonomi, sosial, teknologi, lingkungan, dan hukum, dengan kondisi dan implikasi bagi SIMAK."),
 ("Analisis risiko dan mitigasi", "Tujuh risiko dipetakan menurut dampak dan probabilitas; setiap mitigasi merujuk pada gerbang evaluasi pada naskah.",
  "Peta risiko tiga kali tiga dengan tujuh risiko bernomor, dan tabel risiko beserta mitigasinya."),
 ("Peta jalan dari bukti awal ke pilot dan perluasan", "Jadwal hingga Oktober 2026 mengikuti rencana kerja pada Bab III; tahap sesudahnya adalah rencana lanjutan yang belum berkomitmen.",
  "Peta jalan empat fase dan empat lajur: pembuktian, pipeline dan modul AI, pengguna dan ULD, serta kemitraan dan keberlanjutan."),
]

content = []
content.append(par([("Lampiran F. Profil Nilai Jual dan Analisis Strategis SIMAK", True)], style="Heading1", keep_next=True))
content.append(par([("Lampiran ini menyajikan profil nilai jual SIMAK dan analisis strategis pendukungnya (Gambar F.1 sampai F.13). Seluruh angka bukti berasal dari naskah: audit awal, uji kelayakan awal, dan literatur; penilaian dampak, probabilitas, dan skor bersifat penilaian awal tim. Berkas asli dapat disunting tersedia dalam format PPTX.", False)],
                    jc="both", after=120, size=20))

for i, (title, note, alt) in enumerate(SL, start=1):
    pic_par = OxmlElement("w:p")
    # paragraf gambar: gunakan python-docx untuk membuat inline picture
    from docx.text.paragraph import Paragraph
    p_obj = Paragraph(pic_par, d._body)
    ppr = OxmlElement("w:pPr")
    ppr.append(OxmlElement("w:keepNext"))
    if i > 1:
        ppr.append(OxmlElement("w:pageBreakBefore"))
    sp = OxmlElement("w:spacing"); sp.set(qn("w:before"), "0"); sp.set(qn("w:after"), "0"); ppr.append(sp)
    j = OxmlElement("w:jc"); j.set(qn("w:val"), "center"); ppr.append(j)
    pic_par.append(ppr)
    run = p_obj.add_run()
    shape = run.add_picture(f"{PNG_DIR}/s{i:02d}.png", width=Cm(23.4))
    shape._inline.docPr.set("descr", alt)
    shape._inline.docPr.set("title", f"Gambar F.{i}")
    content.append(pic_par)
    content.append(par([(f"Gambar F.{i}. {title[0].upper() + title[1:]}.", False)], jc="center", before=40, after=40, size=18, keep_next=True))
    content.append(par([(note, False)], jc="both", after=0, size=18, italic=True))

exec(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'lampiran_g_block.py'), encoding='utf8').read())
for p in content:
    final_sect.addprevious(p)
d.save(DST)
print("ok", DST)
