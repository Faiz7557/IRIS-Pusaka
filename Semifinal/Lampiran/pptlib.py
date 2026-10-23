# -*- coding: utf-8 -*-
"""Pustaka kecil python-pptx untuk deck Lampiran F SIMAK (palet diambil dari logo)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
import copy
from lxml import etree

# ---- palet logo
NAVY_D = "031A48"; NAVY = "0B2E63"; INDIGO = "113B8E"; ROYAL = "1C55BB"; BRIGHT = "4680ED"
SKY = "5CA1F7"; AZURE = "277DC8"; TEAL = "24839D"; TEALB = "2C9AAE"
TINT1 = "EAF1FC"; TINT2 = "D6E4FA"; TINT3 = "F4F8FE"; TEAL_L = "DCEFF3"; WHITE = "FFFFFF"
INK = "0B1F4D"; MUTED = "4A5D82"; LINE = "B9CCEC"
FONT = "Times New Roman"


def rgb(h):
    return RGBColor.from_string(h)


def new_prs():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def add_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _apply_run(r, size, bold, italic, color, font=FONT):
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = rgb(color)
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", font)


def fill_tf(tf, paras, size=13, color=INK, bold=False, italic=False, align="l", anchor="t",
            margins=(0.08, 0.05, 0.08, 0.05), spacing=None, line_spacing=None):
    """paras: list of paragraph specs. Spec = str | list of runs | dict(runs=..., align=, bullet=, space_after=)
       run = str | (text, {bold, italic, color, size})."""
    tf.word_wrap = True
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    first = True
    for spec in paras:
        opts = {}
        if isinstance(spec, dict):
            opts = spec; runs = spec["runs"]
        else:
            runs = spec
        if isinstance(runs, str):
            runs = [runs]
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        al = opts.get("align", align)
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "j": PP_ALIGN.JUSTIFY}[al]
        sa = opts.get("space_after", spacing)
        if sa is not None:
            p.space_after = Pt(sa)
        if opts.get("space_before") is not None:
            p.space_before = Pt(opts["space_before"])
        if line_spacing:
            p.line_spacing = line_spacing
        bullet = opts.get("bullet")
        if bullet:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(opts.get("indent", 0.16)))))
            pPr.set("indent", str(-int(Inches(opts.get("indent", 0.16)))))
            bu = etree.SubElement(pPr, qn("a:buFont")); bu.set("typeface", "Arial")
            bc = etree.SubElement(pPr, qn("a:buChar")); bc.set("char", bullet if isinstance(bullet, str) else "•")
        for run in runs:
            if isinstance(run, str):
                t, o = run, {}
            else:
                t, o = run
            r = p.add_run()
            r.text = t
            _apply_run(r, o.get("size", opts.get("size", size)), o.get("bold", opts.get("bold", bold)),
                       o.get("italic", opts.get("italic", italic)), o.get("color", opts.get("color", color)))


def shape(slide, kind, x, y, w, h, fill=None, line=None, lw=0.75, radius=None, shadow=False):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line); shp.line.width = Pt(lw)
    else:
        shp.line.fill.background()
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    if not shadow:
        sp = shp._element.spPr
        eff = etree.SubElement(sp, qn("a:effectLst"))
    shp.text_frame.text = ""
    return shp


def box(slide, x, y, w, h, paras=None, fill=TINT1, line=None, lw=0.75, radius=None, size=13, color=INK, bold=False,
        italic=False, align="l", anchor="t", margins=(0.1, 0.06, 0.1, 0.06), spacing=None, line_spacing=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    s = shape(slide, kind, x, y, w, h, fill=fill, line=line, lw=lw, radius=radius)
    if paras is not None:
        fill_tf(s.text_frame, paras, size=size, color=color, bold=bold, italic=italic, align=align, anchor=anchor,
                margins=margins, spacing=spacing, line_spacing=line_spacing)
    return s


def text(slide, x, y, w, h, paras, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_tf(tb.text_frame, paras, **kw)
    return tb


def line(slide, x1, y1, x2, y2, color=ROYAL, lw=1.25, arrow=False, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(lw)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", "dash")
    if arrow:
        t = etree.SubElement(ln, qn("a:tailEnd")); t.set("type", "triangle"); t.set("w", "med"); t.set("len", "med")
    return c


def circle(slide, cx, cy, d, fill=ROYAL, label=None, size=12, color=WHITE, line_c=None):
    s = shape(slide, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, fill=fill, line=line_c)
    if label is not None:
        fill_tf(s.text_frame, [label], size=size, color=color, bold=True, align="c", anchor="m", margins=(0, 0, 0, 0))
    return s


def chevron(slide, x, y, w, h, fill=ROYAL, paras=None, size=12, color=WHITE, kind=MSO_SHAPE.CHEVRON):
    s = shape(slide, kind, x, y, w, h, fill=fill)
    if paras:
        fill_tf(s.text_frame, paras, size=size, color=color, bold=True, align="c", anchor="m", margins=(0.25, 0.02, 0.15, 0.02))
    return s


def grad_bar(slide, y=0, h=0.13):
    """Batang gradasi navy -> biru -> teal (mengikuti logo)."""
    cols = [NAVY_D, NAVY, INDIGO, ROYAL, BRIGHT, AZURE, TEAL]
    w = 13.333 / len(cols)
    for i, c in enumerate(cols):
        shape(slide, MSO_SHAPE.RECTANGLE, i * w, y, w + 0.01, h, fill=c)


def frame(slide, code, title, subtitle=None, note=None, page=None):
    grad_bar(slide)
    text(slide, 0.5, 0.26, 3.0, 0.3, [[(code, {"bold": True, "color": TEAL, "size": 13})]], margins=(0, 0, 0, 0))
    text(slide, 0.5, 0.5, 12.3, 0.6, [[(title, {"bold": True, "color": NAVY, "size": 25})]], margins=(0, 0, 0, 0))
    if subtitle:
        text(slide, 0.5, 1.02, 12.3, 0.35, [[(subtitle, {"italic": True, "color": MUTED, "size": 13})]], margins=(0, 0, 0, 0))
    line(slide, 0.5, 7.08, 12.83, 7.08, color=LINE, lw=0.75)
    foot = "SIMAK · Innovation Case Competition, Campus Data Week 2026 · Tim IRIS kata HQ ini lomba terakhir"
    text(slide, 0.5, 7.11, 8.6, 0.3, [[(foot, {"color": MUTED, "size": 10})]], margins=(0, 0, 0, 0))
    if note:
        text(slide, 0.5, 6.78, 12.3, 0.3, [[(note, {"italic": True, "color": MUTED, "size": 10.5})]], margins=(0, 0, 0, 0))


def pill(slide, x, y, w, h, label, fill=ROYAL, color=WHITE, size=13, bold=True):
    return box(slide, x, y, w, h, [label], fill=fill, radius=0.5, size=size, color=color, bold=bold, align="c", anchor="m",
               margins=(0.08, 0.02, 0.08, 0.02))


def poly(slide, pts, fill=None, line_c=ROYAL, lw=2.0, alpha=None):
    fb = slide.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]))
    fb.add_line_segments([(Inches(x), Inches(y)) for x, y in pts[1:]], close=True)
    shp = fb.convert_to_shape()
    if fill:
        shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            sf = shp.fill._xPr.find(qn("a:solidFill"))
            clr = sf[0]
            a = etree.SubElement(clr, qn("a:alpha")); a.set("val", str(int(alpha * 1000)))
    else:
        shp.fill.background()
    shp.line.color.rgb = rgb(line_c); shp.line.width = Pt(lw)
    return shp
