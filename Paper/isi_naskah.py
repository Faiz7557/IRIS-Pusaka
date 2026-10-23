"""Mengisi naskah final: teks Bab IV–V, abstrak, revisi Bab I–III dan Lampiran A, D, E sesuai 17 komentar,
rujukan [55]–[56] terverifikasi, lalu menghapus komentar, sorotan, dan penanda ISIAN.

  python isi_naskah.py   (dari folder Final/Paper)

Sumber: salinan rakitan susun_final.py (…_Final_rakitan.docx); keluaran menimpa …_Final.docx.
Markah teks: *miring*, _{subskrip}, ^{superskrip}. Angka dari ANGKA_KANON.md dan lampiran_data.json.
Penanda ISIAN Lampiran B (Turnitin) sengaja dibiarkan: laporan Turnitin final harus dibuat tim.
"""
import hashlib
import re
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "IRIS kata HQ ini lomba terakhir_ICC2026_Final_rakitan.docx"
OUT = ROOT / "IRIS kata HQ ini lomba terakhir_ICC2026_Final.docx"
SIDE = ROOT / ".isian.sha256"
if OUT.exists() and SIDE.exists() and hashlib.sha256(OUT.read_bytes()).hexdigest() != SIDE.read_text().strip():
    raise SystemExit(f"{OUT.name} sudah disunting tangan sejak diisi skrip ini; tidak ditimpa.")

doc = docx.Document(SRC)
body = doc.element.body
TOK = re.compile(r"(\*[^*]+\*|_\{[^}]+\}|\^\{[^}]+\})")


def runs(p, text, size=None, bold=False):
    """Tambahkan teks bermarkah ke paragraf p."""
    for part in TOK.split(text):
        if not part:
            continue
        it = sub = sup = False
        if part.startswith("*"):
            part, it = part[1:-1], True
        elif part.startswith("_{"):
            part, sub = part[2:-1], True
        elif part.startswith("^{"):
            part, sup = part[2:-1], True
        r = p.add_run(part)
        r.italic = it or None
        r.bold = bold or None
        if sub:
            r.font.subscript = True
        if sup:
            r.font.superscript = True
        if size:
            r.font.size = Pt(size)


def kosongkan(p):
    for el in list(p._p):
        if el.tag != qn("w:pPr"):
            p._p.remove(el)


def baru(text, daftar=False):
    """Paragraf badan bergaya semifinal (inden baris pertama, rata kiri-kanan) atau butir daftar bernomor."""
    p = doc.add_paragraph()
    pp = p._p.get_or_add_pPr()
    ind = OxmlElement("w:ind")
    if daftar:
        ind.set(qn("w:left"), "284"); ind.set(qn("w:hanging"), "284")
    else:
        ind.set(qn("w:firstLine"), "142")
    pp.append(ind)
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "both"); pp.append(jc)
    runs(p, text)
    return p


def cari(awal):
    for p in doc.paragraphs:
        if p.text.replace("\xa0", " ").startswith(awal):
            return p
    raise SystemExit(f"tidak ditemukan: {awal!r}")


DAFTAR = re.compile(r"\d+\)\t")


def isi_penanda(kode, paragraf):
    """Ganti penanda [ISIAN TIM — kode …] dengan paragraf. Paragraf penanda dipakai sebagai paragraf terakhir
    karena sebagian membawa sectPr (batas bagian dua kolom → objek selebar halaman)."""
    p = cari(f"[ISIAN TIM — {kode}")
    for t in paragraf[:-1]:
        p._p.addprevious(baru(t, daftar=bool(DAFTAR.match(t)))._p)
    akhir = paragraf[-1]
    kosongkan(p)
    if DAFTAR.match(akhir):
        ind = p._p.pPr.find(qn("w:ind"))
        for a in list(ind.attrib):
            del ind.attrib[a]
        ind.set(qn("w:left"), "284"); ind.set(qn("w:hanging"), "284")
    runs(p, akhir)


def tulis_ulang(p, markah, size=None):
    kosongkan(p)
    runs(p, markah, size=size)


def ke_markah(p):
    """Teks paragraf sebagai markah, agar format miring/subskrip yang sudah ada tidak hilang saat ditulis ulang."""
    out = []
    for r in p.runs:
        t = r.text
        if not t:
            continue
        if r.font.subscript:
            t = "_{" + t + "}"
        elif r.font.superscript:
            t = "^{" + t + "}"
        elif r.italic:
            t = "*" + t + "*"
        out.append(t)
    return "".join(out).replace("**", "")


def ganti(p, lama, baru_):
    """Ganti potongan markah paragraf (atau string); spasi pada `lama` juga cocok dengan spasi tak putus."""
    t = p if isinstance(p, str) else ke_markah(p)
    pola = re.escape(lama).replace(r"\ ", "[ \xa0]")
    t2, n = re.subn(pola, lambda m: baru_, t, count=1)
    assert n, (t[:40], lama)
    return t2


def sel(cell, teks, size=9):
    tulis_ulang(cell.paragraphs[0], teks, size=size)


# =====================================================================================  BAB IV
IV_A = [
    "MVP SIMAK berjalan sebagai aplikasi web lokal (FastAPI) dan dikumpulkan pada repositori "
    "github.com/Faiz7557/IRIS-Pusaka (tag v1.0-final). Gambar 4 memperlihatkan empat halamannya, yang melayani dua "
    "pengguna berbeda. Panel (a), halaman Remediasi & Tinjauan (/), adalah ruang kerja staf ULD. Setelah PDF diunggah, "
    "halaman ini menampilkan katalog gambar beserta jenisnya, hasil veraPDF untuk B0, B1, dan D, serta antrean tinjauan "
    "yang mencantumkan alasan penahanan setiap elemen dengan pilihan setujui, ubah, tolak, atau tandai dekoratif. Pada "
    "dokumen uji materi_03 yang tampil di panel tersebut, jumlah aturan PDF/UA-1 yang gagal turun dari 6 (B0) ke 5 (B1) "
    "dan 0 (D); dua grafik terbit otomatis, tiga ditahan, dan dua gambar disaring sebagai dekoratif. Panel (b), Demo "
    "Gerbang (/demo.html), menjalankan B2 dan SIMAK pada grafik yang sama dan dapat menyuntikkan satu dari empat jenis "
    "kesalahan ke deskripsi SIMAK. Panel (c), halaman Evaluasi (/evaluasi.html), menampilkan isi berkas hasil evaluasi "
    "apa adanya. Panel (d) adalah keluaran HTML yang diterima mahasiswa tunanetra. Setiap gambar membawa penanda status "
    "“Dihasilkan otomatis · angka terverifikasi”, “Disetujui ULD”, atau “Menunggu tinjauan ULD”, dan bilah mode dengar "
    "dapat memutar, menjeda, serta memilih suara. Halaman kelima, Dengar (/dengar.html), membacakan urutan baca PDF asli "
    "dan hasil SIMAK secara berdampingan. Daftar titik akhir layanan dimuat pada Lampiran P.",

    "Pada lapisan model, *T*_{1} dan *T*_{2} diekstraksi oleh Gemma 4 26B-A4B melalui Gemini API dengan suhu 0. Keduanya "
    "memakai model yang sama dengan dua perintah berbeda (hash a2db0023d821 dan bbe13c6a83dc), sehingga perbedaan "
    "keduanya hanya berasal dari perintah, bukan dari model (Bab IV.D). Setiap respons disimpan dalam *cache* dengan kunci "
    "hash citra, perintah, skema ekstraksi (2881acdadcfe), dan model, sehingga evaluasi dapat diulang dan demo dapat "
    "berjalan tanpa jaringan. Mode dengar memakai Web Speech API bawaan peramban dengan suara id-ID yang berjalan lokal "
    "tanpa kunci API; suara ElevenLabs dapat diaktifkan sebagai pilihan. Validasi memakai veraPDF 1.30.3 dengan profil "
    "PDF/UA-1, dan pembanding B1 adalah OpenDataLoader PDF 2.5.11 tanpa modul tambahan. Waktu proses dokumen uji tiga "
    "halaman adalah 29,9–43,9 detik, sedangkan PDF nyata setebal 1–192 halaman memerlukan 6,5–43,0 detik. Bila respons "
    "VLM sudah tersedia di *cache*, materi_03 selesai dalam sekitar 7,6 detik.",

    "Gambar 5 menelusuri satu grafik uji, syn043 (grafik area kunjungan perpustakaan dengan 10 kategori), melalui "
    "gerbang. Kedua ekstraksi identik (*δ* = 0,000 ≤ *θ*_{e} = 0,01), nilai tercetak pada grafik, dan sumbu konsisten. "
    "Kode menghitung maksimum (Jul, 1.452), minimum (Feb, 854), selisih 598, rerata 1.206,8, dan tren naik. Templat "
    "menyusun 16 klaim dari fakta tersebut, dan 16/16 cocok dengan *T*, sehingga deskripsi terbit. Ketika pasangan "
    "label–nilai April dan Oktober ditukar, *c*_{F} gagal dan deskripsi ditahan dengan alasan yang tercatat. B2 pada "
    "grafik yang sama menulis 14 angka. Satu di antaranya, rerata “sekitar 1.107”, tidak setia terhadap nilai acuan "
    "1.206,8, tetapi tetap terbit karena B2 tidak memiliki gerbang. Contoh ini memperlihatkan hal yang tidak tampak pada "
    "keluaran B2: setiap angka D dapat ditelusuri ke tabel data dan ke statistik yang dihitung kode.",

    "Tabel 9 memetakan realisasi sepuluh fitur pada Tabel 6. Enam fitur berjalan penuh (CLS-01, CHT-01, GAT-01, OUT-01, "
    "OUT-02, REV-01) dan dapat ditunjukkan dalam demo. Dua fitur berjalan sebagian. STR-01 memakai pendeteksi tabel "
    "PyMuPDF sebagai cadangan, bukan Docling seperti rancangan, dan tabel cadangan belum ditulis sebagai struktur "
    "Table/TH di PDF; CHT-02 menyusun kalimat murni dari templat tanpa VLM penyunting. Dua fitur rumus belum dibangun: "
    "tanpa pengenal kedua dan MathCAT, setiap rumus ditahan beserta draf LaTeX-nya. Dilihat dari rumusan masalah, STR-01 "
    "dan OUT-01 melayani RM1; CHT-01, CHT-02, dan OUT-02 melayani bagian grafik pada RM2; GAT-01 dan REV-01 melayani "
    "RM3. Bagian rumus pada RM2 belum memiliki fitur yang berjalan.",
]

IV_B = [
    "Evaluasi memakai set uji tertahan berisi 60 grafik sintetis dan 4 dokumen tiga halaman (*seed* 20260930). Set ini "
    "baru dibuka setelah ambang dikunci pada 12 Oktober pukul 12.38 berdasarkan 30 grafik pengembangan (*θ*_{e} = 0,01; "
    "*ε* = 0,02; *ρ* = 0,10; Lampiran M). Versi beku dijalankan dua kali, pada 12 Oktober (*run* 1) dan 23 Oktober "
    "(*run* 2), dengan model, suhu, dan hash perintah yang sama; *run* 2 juga mencatat hash skema yang belum tercatat "
    "pada *run* 1. Tidak ada status grafik yang berubah (0/60) dan seluruh ukuran identik (Lampiran H), sehingga angka "
    "berikut berlaku untuk kedua *run*. Tabel 10 menghadapkan setiap gerbang pada kriteria terkuncinya, Tabel 11 "
    "membandingkan D dengan B0, B1, dan B2, Gambar 6 memperlihatkan status 60 grafik pada B2 dan D, dan Gambar 7 "
    "menempatkan selang Wilson 95% (8) terhadap garis kriteria. Proporsi ditulis sebagai *k*/*n* [batas bawah; batas "
    "atas].",

    "*G1 (struktur dan metadata).* D lulus G1 pada 4/4 dokumen [0,510; 1,000], sedangkan B0 dan B1 sama-sama 0/4 "
    "[0,000; 0,490]. Pada setiap dokumen, aturan PDF/UA-1 yang gagal turun dari 6 (B0) ke 5 (B1) dan 0 (D). Kriteria "
    "terpenuhi pada proporsi teramati, tetapi dengan *n* = 4 batas bawah selangnya hanya 0,510, sehingga data ini belum "
    "cukup untuk menyatakan bahwa proporsi sebenarnya mencapai 0,90. Sesuai definisi pada Bab III.G, G1 menghitung "
    "penanda “[Menunggu tinjauan ULD]” pada elemen yang ditahan sebagai teks alternatif yang ada. Bila penanda itu "
    "dihitung belum bermakna (G1-ketat, pelaporan tambahan), tidak ada dokumen D yang lulus, 0/4 [0,000; 0,490], "
    "karena setiap dokumen uji memuat sedikitnya satu gambar yang ditahan.",

    "*G3 (kesetiaan angka).* D menerbitkan 40/60 grafik [0,541; 0,773] dengan 440 klaim (8–16 klaim per grafik). "
    "Seluruh klaim itu setia terhadap tabel acuan, *F*_{k} = 440/440 [0,991; 1,000], sehingga kriteria *F*_{k} ≥ 0,98 "
    "terpenuhi. Pada uji injeksi, satu kesalahan disuntikkan ke deskripsi setiap grafik yang terbit untuk setiap jenis "
    "yang berlaku, dan 137/138 [0,960; 0,999] deskripsi terinjeksi tertahan (kriteria ≥ 0,90). Menurut jenisnya, ubah "
    "angka tertahan 40/40, tukar label 39/40, turunan salah 40/40, dan tren salah 18/18; injeksi tren hanya berlaku pada "
    "grafik berurutan. Satu-satunya yang lolos adalah tukar label pada syn073. Kedua nilai yang ditukar berselisih 5, di "
    "bawah toleransi *ε*·*s* = 6,74, sehingga menurut (5) kedua klaim masih tergolong setia (Lampiran K).",

    "Rincian per grafik (Lampiran J) memperlihatkan pola penahanan yang tegas. Keempat puluh grafik yang terbit memiliki "
    "label nilai, sedangkan kedua puluh grafik yang ditahan tidak. Sembilan belas grafik ditahan hanya oleh *c*_{L}, dan "
    "satu grafik (syn014) ditahan oleh *c*_{L} dan *c*_{δ} (*δ* = 0,0128 > *θ*_{e}). Kelima grafik dengan *δ* > 0 "
    "seluruhnya ditahan. Bila klaim D pada grafik yang ditahan dibandingkan dengan acuan, hanya 205/222 [0,881; 0,952] "
    "yang setia, dan 6/20 grafik memuat sedikitnya satu klaim tidak setia; yang terburuk, syn068, hanya 1/9.",

    "*Pembanding B2.* Tanpa gerbang, B2 menerbitkan 60/60 grafik. Pemeriksa pola menandai 17 angka pada 7/60 grafik "
    "[0,058; 0,222] sebagai tidak setia, sedangkan D tidak menerbitkan satu pun angka keliru. *F*_{k} B2 diperkirakan "
    "566/583 = 0,971 [0,954; 0,982]; nilai ini perkiraan karena teks bebas diperiksa dengan pola, bukan dengan klaim "
    "terstruktur. Audit manual atas ketujuh belas tanda (Lampiran L) menemukan 13 pembulatan atau perkiraan di luar *ε* "
    "yang seluruhnya terjadi pada grafik tanpa label nilai, 3 rerata yang dihitung model secara keliru (syn043, syn083, "
    "syn068), dan 1 salah tandai oleh pemeriksa (syn079), sehingga 16/17 tanda benar. Dari tujuh grafik bertanda, lima "
    "tidak berlabel nilai. Pada dua grafik berlabel, kesalahan nyata hanya rerata syn043, sedangkan tanda pada syn079 "
    "keliru.",

    "*G4 (kesetaraan informasi).* Dua puluh grafik masing-masing diberi empat pertanyaan yang dibangkitkan dari tabel "
    "data. Penjawab yang hanya membaca deskripsi D menjawab benar 80/80 [0,954; 1,000], dari deskripsi B2 79/80 "
    "[0,932; 0,998], dan dari tabel data 80/80. Menurut (10), *Q*(D) = 1,00 dan *Q*(B2) = 0,99, sehingga keduanya "
    "melewati kriteria *Q* ≥ 0,80.",

    "*G6 ringan dan gerbang lain.* Pada enam PDF nyata dari domain ac.id (Common Crawl CC-MAIN-2026-39; 1–192 halaman), "
    "D lulus G1 pada 6/6 dokumen [0,610; 1,000], sedangkan B0 dan B1 0/6 [0,000; 0,390]. Kelulusan veraPDF penuh hanya "
    "tercapai pada 2/6 dokumen [0,097; 0,700]; empat dokumen lain masih gagal pada satu sampai enam aturan. Seluruh 18 "
    "elemen ditahan, *R* = 18/18 [0,824; 1,000], dan tidak ada grafik yang terbit; 15/18 elemen tersebut adalah figur "
    "vektor (Lampiran N). Karena baru mencakup 6 dari 60 dokumen dan *κ* serta pemeriksaan NVDA belum dilakukan, G6 "
    "dinyatakan sebagian. G2, G5, dan G7 tidak dijalankan. Dengan demikian, tidak ada kriteria yang gagal di antara "
    "gerbang yang dijalankan, tetapi empat gerbang belum lengkap atau belum dijalankan.",
]

IV_C = [
    "*Rumusan masalah 1: terjawab sebagian.* Perakitan keluaran SIMAK menutup seluruh aturan PDF/UA-1 yang dapat "
    "diperiksa mesin pada dokumen uji dan menambahkan tiga hal yang paling sering hilang pada audit awal, yaitu bahasa "
    "dokumen, judul, dan baris judul tabel. Selisih antara B1 dan D juga menunjukkan bahwa *auto-tag* saja tidak cukup: "
    "B1 hanya mengurangi satu aturan dan gagal G1 pada semua dokumen. Namun, lolos G1 belum berarti terbaca utuh oleh "
    "pembaca layar. G1 memeriksa keberadaan teks alternatif, bahasa, dan judul, bukan urutan baca. G2, yang mengukur "
    "urutan baca dengan *τ* pada (1), tidak dijalankan, sehingga urutan baca D sepenuhnya bergantung pada urutan blok "
    "OpenDataLoader yang belum diverifikasi. Pada dokumen nyata, kelulusan veraPDF penuh hanya 2/6, tabel cadangan belum "
    "memiliki sel TH, dan G1-ketat 0/4 menunjukkan bahwa dokumen uji yang lulus G1 masih memuat gambar yang menunggu "
    "tinjauan. Jawaban atas RM1 karenanya terbatas pada metadata dan teks alternatif pada sampel kecil; urutan baca "
    "belum terbukti.",

    "*Rumusan masalah 2: terjawab sebagian untuk grafik, belum terjawab untuk rumus.* Pada set uji, tidak ada klaim "
    "keliru di antara 440 klaim sintetis yang terbit, sedangkan B2 dari model yang sama menerbitkan 17 angka bertanda "
    "pada 7/60 grafik. Audit B2 menunjukkan letak kesalahan yang sebenarnya. Pada grafik berlabel nilai, model membaca "
    "angka dengan tepat, dan kesalahan muncul ketika model menghitung sendiri statistik turunan, seperti rerata syn043 "
    "yang ditulis sekitar 1.107 padahal 1.206,8. Statistik yang dihitung kode menghilangkan jenis kesalahan ini. Pada "
    "grafik tanpa label nilai, kedua ekstraksi D memperkirakan nilai dari skala dengan kesalahan yang sama sehingga *δ* "
    "tetap nol, sesuai peringatan Bab II.C bahwa ekstraksi ganda hanya menangkap kesalahan acak. Tanpa *c*_{L}, 17 klaim "
    "tidak setia pada grafik yang ditahan (205/222 setia) akan sampai ke pembaca. Kesetiaan keluaran D karena itu "
    "berasal dari statistik yang dihitung kode dan dari keputusan untuk tidak menerbitkan grafik yang tidak dapat "
    "diverifikasi; VLM-nya tidak lebih akurat daripada B2. Batasnya tetap penting. Set ujinya sintetis. *T*_{1} dan "
    "*T*_{2} berasal dari satu model, sehingga *c*_{F} terhadap *T*_{1} bernilai 1 menurut konstruksi dan kesetiaan "
    "harus diukur terpisah terhadap *T*_{acuan}. *Q* juga hampir jenuh karena pertanyaan dibangkitkan dari tabel dan "
    "dijawab oleh model yang sama, sehingga G4 hanya membedakan D dan B2 sebesar satu jawaban. Bagian rumus pada RM2 "
    "belum terjawab karena setiap rumus ditahan.",

    "*Rumusan masalah 3: terjawab untuk grafik pada set uji.* Gerbang menahan setiap grafik yang tidak memenuhi (7) dan "
    "menyerahkannya ke antrean ULD dengan alasan per elemen: 20/60 grafik uji ditahan, dan pada dokumen nyata seluruh 18 "
    "elemen masuk antrean. Penahanan ini beralasan, karena 6/20 grafik yang ditahan memang memuat klaim keliru. Uji "
    "injeksi 137/138 menunjukkan bahwa logika gerbang menahan kesalahan di atas toleransi yang disuntikkan ke deskripsi, "
    "tetapi tidak membuktikan bahwa VLM membaca grafik dengan benar. Kesalahan baca VLM ditangani oleh *c*_{δ}, *c*_{L}, "
    "dan *c*_{A}, sedangkan kesalahan di bawah *ε*·*s* tidak terdeteksi menurut definisi, seperti pada syn073. Harga "
    "*fail-closed* juga terlihat. Empat belas grafik yang ditahan sebenarnya benar seluruh klaimnya, dan pada dokumen "
    "nyata *R* bernilai 18/18, sehingga untuk materi seperti sampel G6 SIMAK belum mengurangi jumlah elemen yang harus "
    "ditinjau ULD. Yang berubah adalah sifat pekerjaannya: staf meninjau draf beserta alasan penahanannya dan tidak "
    "menulis teks alternatif dari nol, sedangkan mahasiswa melihat status setiap gambar alih-alih teks pengisi.",

    "*Kaitan dengan pustaka.* Temuan ini sejalan dengan CharXiv [15], yang melaporkan bahwa VLM masih jauh tertinggal "
    "dari manusia dalam memahami grafik. Pada set uji, model gagal ketika harus memperkirakan nilai dari skala dan "
    "ketika menghitung statistik sendiri. SIMAK tidak memperbaiki kemampuan model itu; SIMAK membatasi dampaknya dengan "
    "hanya menerbitkan grafik yang nilainya dapat diperiksa. Kumar dkk. [30] mendapati LLM lemah dalam menilai mutu "
    "teks alternatif dan menyarankan pendekatan hibrida. B1 menguatkan temuan itu dari sisi validator, karena teks "
    "pengisi “image N” tetap lolos veraPDF (Tabel 8) tetapi ditolak G1. SIMAK menjalankan bentuk hibrida tersebut: "
    "aturan veraPDF untuk struktur, model untuk ekstraksi, kode untuk statistik, dan ULD untuk elemen yang tidak dapat "
    "diverifikasi. Deskripsi yang dibangun dari tabel data juga mengikuti saran Jung dkk. [12] agar teks alternatif "
    "merujuk pada data, dan mencakup level 1–3 model Lundgard dan Satyanarayan [26]. Manfaat deskripsi tersebut bagi "
    "pembaca tunanetra belum diuji (G7).",
]

IV_D = [
    "Bagian yang belum berjalan atau belum sempat diuji dirinci berikut ini beserta buktinya.",
    "1)\tSet uji G1, G3, dan G4 adalah data sintetis buatan tim, yaitu 60 grafik dan 4 dokumen, jauh di bawah rancangan "
    "150 grafik, 30 dokumen, dan 200 rumus; *benchmark* publik (VisText, ChartQA, UniMER-Test, IndoTabVQA) tidak dipakai "
    "(Lampiran O).",
    "2)\tG6 ringan tidak berisi materi kuliah. Enam PDF yang lolos aturan URL ternyata dua RPS (ITS), dua *module "
    "handbook* (Unair, Unej), serta SOP cuti dan pengumuman (UNY). Dokumen ini tidak memuat grafik raster, dan presisi "
    "aturan URL belum diperiksa.",
    "3)\tFigur vektor belum dideskripsikan. Sebanyak 15/18 elemen antrean G6 adalah figur vektor yang ditahan tanpa "
    "deskripsi dan belum tampil sebagai kartu di antrean antarmuka; 13 di antaranya berasal dari satu dokumen (cc_03).",
    "4)\t*T*_{1}, *T*_{2}, dan penjawab G4 memakai satu model (Gemma 4 26B-A4B), sehingga kesalahan sistematiknya dapat "
    "sama. Model *T*_{2} dapat diganti melalui variabel SIMAK_MODEL_T2, tetapi belum diuji.",
    "5)\t*c*_{F} terhadap *T*_{1} bernilai 1 menurut konstruksi karena deskripsi dibangun dari *T*_{1}; perlindungan "
    "nyata berasal dari *c*_{δ}, *c*_{L}, dan *c*_{A}. Uji injeksi menguji logika gerbang, bukan kesalahan baca VLM, "
    "sehingga kesetiaan diukur terpisah terhadap *T*_{acuan}.",
    "6)\tG1 menghitung penanda tahan sebagai teks alternatif yang ada, sehingga hasil G1 (4/4) dan G1-ketat (0/4) "
    "berbeda jauh.",
    "7)\tModul rumus belum dibangun dan semua rumus ditahan. Gambar dekoratif belum ditandai sebagai Artifact pada PDF, "
    "dan tabel cadangan belum ditulis sebagai Table/TH.",
    "8)\tG2, G5, dan G7 tidak dijalankan, dan G6 penuh (60 dokumen, *κ* dua penilai, pemeriksaan NVDA) belum dilakukan. "
    "Belum ada uji bersama mahasiswa tunanetra maupun staf ULD.",
    "9)\tSkema ekstraksi berubah setelah *run* 1, dan hash skema tidak tercatat pada manifes *run* 1. *Run* 2 dijalankan "
    "pada kode akhir, memanggil ulang VLM untuk grafik yang respons *cache*-nya memakai skema lama, dan menghasilkan "
    "ukuran yang identik.",
    "10)\tToleransi *ε* membuat kesalahan di bawah *ε*·*s* tidak terdeteksi menurut definisi, seperti tukar label "
    "syn073 (5 < 6,74). Nilai *ε* = 0,02 dipilih pada set pengembangan, dan dampaknya bagi pembaca belum diuji.",
]

# =====================================================================================  BAB V
V_A = [
    "Penelitian ini membangun MVP SIMAK dan mengujinya pada set uji tertahan dengan ambang yang dikunci lebih dahulu. "
    "RM1 terjawab sebagian: SIMAK melengkapi teks alternatif, bahasa, judul, dan baris judul tabel sehingga keempat "
    "dokumen uji lulus G1 tanpa kegagalan veraPDF, berbeda dengan B0 dan B1, tetapi urutan baca belum diukur dan "
    "dokumen nyata baru lulus veraPDF penuh pada 2/6. RM2 terjawab untuk grafik pada data sintetis: 440/440 klaim yang "
    "terbit setia terhadap tabel acuan, sedangkan VLM yang sama tanpa gerbang menerbitkan 17 angka bertanda pada 7/60 "
    "grafik. Kesetiaan itu dicapai dengan menghitung statistik di kode dan tidak menerbitkan grafik yang nilainya tidak "
    "tercetak. Bagian rumus pada RM2 belum terjawab karena semua rumus ditahan. RM3 terjawab untuk grafik pada set uji: "
    "elemen yang tidak terverifikasi ditahan beserta alasannya (*R* = 20/60 pada set uji dan 18/18 pada dokumen nyata), "
    "dan 137/138 deskripsi terinjeksi tertahan. Hasil ini menunjukkan bahwa gerbang mencegah angka keliru terbit tanpa "
    "penanda pada set uji. Kinerja pada materi kuliah nyata yang memuat grafik, manfaat bagi pembaca tunanetra, dan "
    "beban kerja ULD belum terbukti.",
]

V_B = [
    "Urutan langkah lanjutan disusun menurut sumber beban tinjauan pada G6. Pertama, ekstraksi data dari jalur vektor "
    "PDF, karena figur vektor menyumbang sebagian besar antrean dokumen nyata dan datanya dapat dibaca langsung tanpa "
    "perkiraan skala. Kedua, *T*_{2} dengan model berbeda (misalnya gemma-4-31b-it melalui SIMAK_MODEL_T2) agar "
    "kesalahan sistematik kedua ekstraksi tidak lagi sama, dengan biaya sekitar 40–70 detik per panggilan. Ketiga, "
    "pengenal rumus kedua dan MathCAT dengan aturan bahasa Indonesia untuk menyelesaikan FRM-01 dan FRM-02. Keempat, "
    "penandaan Artifact untuk gambar dekoratif dan struktur Table/TH untuk tabel cadangan.",
    "Pada sisi evaluasi, G6 penuh dijalankan pada 60 materi kuliah nyata, dengan 28 tautan audit awal sebagai kandidat, "
    "dua penilai yang kesepakatannya diukur dengan *κ*, dan pemeriksaan NVDA; G2 dan G5 menyusul pada data berlabel. "
    "G7 dijalankan bersama ULD mitra untuk mengukur ketepatan dan waktu tugas mahasiswa tunanetra serta waktu kerja "
    "staf. Audit 600 PDF dari bingkai Common Crawl dan pemakaian *benchmark* publik (VisText, ChartQA, UniMER-Test, "
    "IndoTabVQA) diperlukan agar kesimpulan tidak bergantung pada data buatan tim.",
]

V_C = [
    "Bila dipakai, SIMAK memberi mahasiswa tunanetra kepastian tentang status setiap gambar. Deskripsi yang terbit "
    "dapat dipercaya angkanya, sedangkan gambar yang belum diverifikasi diberi penanda alih-alih teks pengisi yang "
    "menyesatkan. Pada set uji sintetis tidak ada angka keliru yang terbit, dan mode dengar id-ID berjalan lokal tanpa "
    "biaya layanan. Pengaruhnya pada ketepatan dan kecepatan mahasiswa menyelesaikan tugas belum diukur dan menjadi "
    "sasaran G7. Bagi staf ULD, SIMAK mengubah pekerjaan dari menulis teks alternatif menjadi meninjau draf yang "
    "disertai alasan penahanan, dan beban tinjauannya terukur sejak awal melalui *R*; penghematan waktu kerja masih "
    "hipotesis. Bagi dosen, materi dapat diunggah apa adanya tanpa menulis teks alternatif sendiri, meskipun adopsinya "
    "di kampus belum diuji. Mahasiswa yang memilih mendengarkan materi dapat memakai mode dengar yang sama, tetapi "
    "manfaat belajarnya tetap hipotesis seperti dinyatakan pada Bab I. Bagi perguruan tinggi, SIMAK membantu ULD "
    "menjalankan kewajiban akomodasi yang layak [2, 3] dan mendukung SDG 4 dan 10.",
]

for kode, isi in (("IV.A", IV_A), ("IV.B", IV_B), ("IV.C", IV_C), ("IV.D", IV_D),
                  ("V.A", V_A), ("V.B", V_B), ("V.C", V_C)):
    isi_penanda(kode, isi)

# =====================================================================================  ABSTRAK
ABSTRAK = (
    "— Materi kuliah PDF jarang dapat dibaca utuh oleh pembaca layar. Audit awal atas 28 PDF dari 18 kampus menemukan "
    "0/28 dokumen memenuhi PDF/UA-1, dan hanya 29/1.111 elemen gambar pada dokumen bertag memiliki teks alternatif. "
    "Makalah ini menyajikan SIMAK, *pipeline* remediasi yang melengkapi struktur PDF, menyusun deskripsi grafik dari "
    "tabel data yang diekstraksi dua kali oleh model *vision-language* dengan statistik yang dihitung kode, menyediakan "
    "mode dengar berbahasa Indonesia, dan menahan setiap keluaran yang tidak dapat diverifikasi untuk ditinjau Unit "
    "Layanan Disabilitas (ULD) melalui gerbang *fail-closed*. MVP dievaluasi pada set uji tertahan berisi 60 grafik dan "
    "4 dokumen sintetis dengan ambang yang dikunci sebelumnya, dalam dua *run* yang hasilnya identik. Semua dokumen uji "
    "lulus pemeriksaan struktur dan metadata (4/4), sedangkan *auto-tag* sumber terbuka 0/4. Seluruh 440 klaim pada 40 "
    "deskripsi grafik yang terbit setia terhadap data acuan (440/440), sedangkan model yang sama tanpa gerbang "
    "menerbitkan 17 angka yang ditandai keliru pada 7/60 grafik. Sebanyak 137/138 deskripsi yang sengaja diberi kesalahan tertahan. "
    "Harga penahanan adalah beban tinjauan ULD: 20/60 grafik uji dan 18/18 elemen pada enam PDF nyata masuk antrean. "
    "Hasil ini terbatas pada data sintetis; pembacaan rumus, urutan baca, dan uji bersama mahasiswa tunanetra belum "
    "dilakukan."
)
pa = cari("Abstrak—")
kosongkan(pa)
r = pa.add_run("Abstrak")
r.bold = r.italic = True
r.font.size = Pt(9)
runs(pa, ABSTRAK, size=9, bold=True)

# =====================================================================================  REVISI BAB I–III
T3 = next(t for t in doc.tables if t.rows[0].cells[0].text.startswith("Tabel 3"))
status3 = [
    "G1 lulus 4/4 (uji) dan 6/6 (nyata); G1-ketat 0/4; G6 penuh belum (Bab IV.B)",
    "*F*_{k} 440/440 pada klaim terbit; injeksi 137/138 tertahan; set uji sintetis",
    "Grafik: *Q* = 1,00 (G4); rumus belum (G5 tidak dijalankan); manfaat bagi pendengar tetap hipotesis",
    "veraPDF gagal 6 → 5 → 0 (B0, B1, D) pada set uji; lulus penuh 2/6 PDF nyata; urutan baca (G2) belum diukur",
    "Antrean tinjauan berjalan; *R* = 20/60 (uji) dan 18/18 (nyata); dokumen tidak disimpan; G7 tidak dijalankan",
]
for row, s in zip(T3.rows[2:], status3):
    sel(row.cells[2], s)
sel(T3.rows[0].cells[0], "Tabel 3. Manfaat yang dijanjikan, ukuran pembuktinya, dan status buktinya pada MVP")
p = cari("Berdasarkan rumusan masalah tersebut")
tulis_ulang(p, ganti(p, "status buktinya hari ini", "status buktinya pada MVP"))

# II.C: cadangan tabel yang benar-benar dipakai
p = cari("Struktur dokumen dan urutan baca.")
tulis_ulang(p, ganti(p,
    "Penelitian ini memakai OpenDataLoader sebagai tulang punggung karena telah menghasilkan PDF bertag, dan Docling "
    "hanya bila tabel tanpa garis tidak terbaca pada set pengembangan (citra tabel IndoTabVQA dan dokumen kuliah).",
    "Penelitian ini memakai OpenDataLoader sebagai tulang punggung karena telah menghasilkan PDF bertag. Untuk tabel "
    "tanpa garis yang tidak terbaca, MVP memakai pendeteksi tabel PyMuPDF (find_tables) sebagai cadangan, sedangkan "
    "Docling tetap menjadi alternatif yang belum diuji."))

# II.C pers. (7): cL
p = cari("Gerbang verifikasi. Setiap elemen")
tulis_ulang(p,
    "*Gerbang verifikasi.* Setiap elemen *e* memiliki himpunan pemeriksaan *K*(*e*) menurut jenisnya: untuk grafik, "
    "*c*_{F} (*F*_{k} = 1), *c*_{δ} (*δ* ≤ *θ*_{e}), *c*_{A} (rentang sumbu), dan *c*_{L} (nilai tercetak pada grafik); "
    "untuk rumus, *c*_{R} (*A* ≥ *θ*_{r}) antara dua pengenal; foto dan diagram nondata tidak memiliki pemeriksaan "
    "sehingga *K*(*e*) = ∅. Elemen terbit otomatis hanya jika")
p = cari("dengan 1[·] fungsi indikator")
tulis_ulang(p, ke_markah(p) +
    " Pemeriksaan *c*_{L} ditambahkan dari analisis kesalahan pada set pengembangan sebelum ambang dikunci. Pada "
    "kalibrasi versi pertama, *F*_{k} keluaran terbit hanya 0,937 walaupun *δ* = 0, dan grafik tanpa label nilai "
    "memiliki *F*_{k} 0,75–0,875, karena kedua ekstraksi memperkirakan nilai dari skala dengan kesalahan yang sama "
    "(Lampiran M).")

# III.B Arsitektur
p = cari("SIMAK dibangun sebagai aplikasi web tiga lapis")
t = ganti(p,
    "Lapisan model berupa adaptor VLM yang tidak terikat penyedia sehingga model komersial dapat diganti model terbuka lokal.",
    "Lapisan model berupa adaptor VLM yang tidak terikat penyedia sehingga model komersial dapat diganti model terbuka "
    "lokal. MVP memakai Gemma 4 26B-A4B melalui Gemini API dengan suhu 0, dan setiap respons disimpan dalam *cache* "
    "menurut hash citra, perintah, skema, dan model. Karena keterbatasan kuota, *T*_{1} dan *T*_{2} memakai model yang "
    "sama dengan dua perintah berbeda; model *T*_{2} dapat diganti melalui variabel SIMAK_MODEL_T2.")
t = ganti(t, "Keluaran HTML memuat mode dengar berbasis Web Speech API; suara bahasa Indonesia (Microsoft Andika, id-ID) "
          "terpasang pada sistem uji, sedangkan ketersediaannya pada peramban pengguna diuji pada tahap implementasi.",
    "Keluaran HTML memuat mode dengar berbasis Web Speech API dengan suara id-ID bawaan sistem (misalnya Microsoft "
    "Andika) yang berjalan lokal tanpa kunci API; suara ElevenLabs dapat diaktifkan sebagai pilihan.")
tulis_ulang(p, t)

# III.C dua run
p = cari("Gambar 2 memperlihatkan alur penelitian.")
tulis_ulang(p, ke_markah(p) +
    " Janji ini dipenuhi oleh dua *run*, pada 12 dan 23 Oktober 2026, yang menghasilkan status identik pada 60 grafik "
    "uji (Bab IV.B; Lampiran H).")

# III.D Tabel 5 tetap rancangan
p = cari("Tabel 5 merinci data yang dipakai")
tulis_ulang(p, ganti(p, "Tabel 5 merinci data yang dipakai;", "Tabel 5 merinci rancangan data;") +
    " Realisasi data pada MVP lebih kecil daripada rancangan ini dan dirinci pada Lampiran O.")

# Tabel 6: kolom Model/metode dan status rumus
T6 = next(t for t in doc.tables if t.rows[0].cells[0].text.startswith("Tabel 6"))
for row in T6.rows[2:]:
    fid = row.cells[0].text.strip()
    if fid == "STR-01":
        sel(row.cells[2], "OpenDataLoader [35]; PyMuPDF find_tables sebagai cadangan tabel")
    elif fid == "CHT-02":
        sel(row.cells[2], "Kode: maks, min, selisih, rerata, regresi (3), (4); templat bahasa Indonesia")
    elif fid in ("FRM-01", "FRM-02"):
        sel(row.cells[4], "Belum diimplementasikan; rumus ditahan (*fail-closed*)")
p = cari("Setiap fitur AI dikembangkan dengan kerangka tujuh langkah")
tulis_ulang(p, ganti(p, "fitur yang belum dibangun dinyatakan sebagai rencana.",
                     "statusnya mengikuti MVP final, dan realisasinya dirinci pada Tabel 9."))

# III.F modul grafik dan rumus
p = cari("Modul grafik.")
tulis_ulang(p, ganti(p,
    "Kalimat dibangun dari templat bahasa Indonesia yang diisi fakta tersebut, dengan VLM hanya berwenang memilih urutan "
    "dan menyunting kelancaran bahasa tanpa mengubah token angka dan label, sehingga ruang model untuk mengarang angka "
    "dipersempit sejak awal.",
    "Kalimat dibangun sepenuhnya dari templat bahasa Indonesia yang diisi fakta tersebut. VLM tidak menulis kalimat, "
    "sehingga model tidak memiliki ruang untuk mengarang angka pada deskripsi."))
p = cari("Modul rumus.")
tulis_ulang(p,
    "*Modul rumus.* Rancangannya mengubah citra rumus menjadi LaTeX oleh dua pengenal independen, kemudian menjadi "
    "MathML yang diucapkan MathCAT dengan aturan bahasa Indonesia. Karena MathCAT sudah mendukung bahasa Indonesia [43], "
    "kontribusi penelitian ini berada pada jalur dari PDF menuju MathML yang benar, bukan pada aturan ucapannya; "
    "kebenaran jalur itu diperiksa oleh kesepakatan kedua pengenal menurut (6). Pada MVP, pengenal kedua dan MathCAT "
    "belum terpasang, sehingga seluruh rumus masuk antrean tinjauan beserta draf LaTeX dari VLM.")

# III.G G1-ketat sebagai pelaporan tambahan
p = cari("Sistem D dibandingkan dengan tiga pembanding")
tulis_ulang(p, ganti(p, "atau kata tunggal seperti “gambar”).",
    "atau kata tunggal seperti “gambar”). Sebagai pelaporan tambahan, bukan kriteria baru, G1-ketat juga menghitung "
    "penanda tahan “[Menunggu tinjauan ULD]” sebagai teks alternatif yang belum bermakna."))

# III.H dan III.I: jadwal yang sudah lewat
p = cari("Dua hal belum teruji")
tulis_ulang(p,
    "Dua hal belum teruji pada tahap ini, yaitu akurasi ekstraksi tabel data dari grafik oleh VLM dan kualitas keluaran "
    "pada materi kuliah asli. Keduanya diuji pada tahap pengembangan MVP, dan hasilnya dilaporkan pada Bab IV.")
p = cari("Rencana kerja (Lampiran F.13)")
tulis_ulang(p,
    "Rencana kerja (Lampiran F.13) dijalankan dengan titik keputusan lebih awal. Kalibrasi pada 30 grafik pengembangan "
    "menunjukkan bahwa nilai yang tidak tercetak tidak dapat diverifikasi oleh ekstraksi ganda, sehingga *c*_{L} "
    "ditambahkan sebelum ambang dikunci pada 12 Oktober (Lampiran M); set uji tertahan dievaluasi pada hari yang sama "
    "dan diulang pada 23 Oktober. Lingkup minimum demo terpenuhi, yaitu unggah PDF, HTML aksesibel dan PDF bertag, "
    "deskripsi grafik dengan gerbang, antrean tinjauan, dan mode dengar; respons VLM disimpan dalam *cache* agar demo "
    "dapat berjalan luring. Ditunda: PowerPoint dan Word langsung, braille, modul rumus, dan G7; rumus ditahan ke "
    "antrean tinjauan (*fail-closed*).")

# =====================================================================================  LAMPIRAN A, D, E
p = cari("Tim menggunakan Claude")
tulis_ulang(p,
    "Tim menggunakan Claude (Anthropic) melalui Claude Code. Pada babak semifinal, Claude dipakai untuk penelusuran "
    "awal pustaka, pemeriksaan angka terhadap sumber, pembuatan skrip audit dokumen dan skrip uji kelayakan awal, "
    "penyusunan dan revisi draf naskah, serta pemeriksaan bahasa. Pada babak final, Claude dipakai untuk: (1) membantu "
    "menulis dan memperbaiki kode prototipe SIMAK beserta README-nya; (2) membuat skrip pengolah hasil evaluasi dan "
    "penghitung selang Wilson, Tabel 9–11, tangkapan layar Gambar 4, Gambar 5–7, grafik lampiran, serta Lampiran H–P; "
    "(3) merakit struktur dokumen dan membuat alat pemeriksa angka naskah; (4) menelusuri dan memeriksa rujukan [55] "
    "dan [56]; (5) menyusun draf teks Bab IV, Bab V, dan abstrak, serta revisi Bab I–III dan Lampiran D–E, dari hasil "
    "evaluasi dan arahan tim; dan (6) menyusun rencana kerja dan audit internal. Gemma 4 26B-A4B (Google) melalui Gemini "
    "API adalah komponen sistem SIMAK yang dievaluasi, bukan alat penulisan; ElevenLabs adalah komponen suara opsional. "
    "Seluruh isi telah ditinjau dan disunting oleh penulis, dan penulis bertanggung jawab penuh atas akurasi dan "
    "orisinalitas karya. Seluruh referensi telah diperiksa keasliannya oleh penulis.")

p = cari("Perangkat: laptop tanpa GPU")
tulis_ulang(p, ganti(p, "veraPDF 1.30.2 (profil PDF/UA-1).",
                     "veraPDF 1.30.2 (profil PDF/UA-1); MVP final memakai veraPDF 1.30.3 (Lampiran H)."))

p = cari("Antarmuka: HTML, CSS, dan JavaScript")
tulis_ulang(p,
    "Antarmuka: HTML, CSS, dan JavaScript tanpa kerangka kerja berat, dengan halaman Remediasi & Tinjauan, Demo "
    "Gerbang, Evaluasi, Dengar, dan keluaran HTML. Layanan: Python 3.13 dengan FastAPI dan antrean kerja dalam proses; "
    "sesi disimpan pada direktori sementara dan dapat dihapus. Pemroses dokumen: OpenDataLoader PDF 2.5.11 (Java 21), "
    "PyMuPDF 1.28, dan pikepdf 10.16; validasi memakai veraPDF 1.30.3. Model: Gemma 4 26B-A4B melalui Gemini API "
    "(pustaka google-genai) dengan suhu 0 dan *cache* per hash; adaptor VLM menerima citra dan skema JSON lalu "
    "mengembalikan tabel data, sehingga model terbuka (misalnya Qwen2.5-VL) dapat dipasang. Rumus: MathCAT melalui "
    "WebAssembly direncanakan tetapi belum diintegrasikan, sehingga rumus ditahan. Mode dengar: Web Speech API dengan "
    "suara id-ID; ElevenLabs opsional.")
p = cari("Rekaman elemen memuat id")
t = re.sub(r"Titik akhir yang direncanakan:.*$",
           "Titik akhir yang terimplementasi dirinci pada Tabel P.1 (Lampiran P).", ke_markah(p))
tulis_ulang(p, ganti(t, "*c*_{A}, *c*_{R}", "*c*_{A}, *c*_{L}, *c*_{R}"))
p = cari("Sumber daya yang terukur pada uji awal")
tulis_ulang(p, ganti(p, "Waktu dan biaya VLM per gambar diukur pada uji penentu 1–3 Oktober.",
                     "Waktu proses MVP per dokumen dilaporkan pada Bab IV.A."))

# =====================================================================================  REFERENSI [55], [56]
REF = {
    "[55]": "NASA Jet Propulsion Laboratory dan PDF Association (program DARPA SafeDocs), “SAFEDOCS "
            "(CC-MAIN-2021-31-PDF-UNTRUNCATED) corpus,” Digital Corpora, 2021. Diakses: 23 Okt. 2026. Tersedia: "
            "digitalcorpora.org/corpora/file-corpora/cc-main-2021-31-pdf-untruncated/",
    "[56]": "A. Rossberg, Ed., “WebAssembly core specification,” W3C Recommendation, 5 Des. 2019. Diakses: 23 Okt. "
            "2026. Tersedia: www.w3.org/TR/wasm-core-1/",
}
for no, teks in REF.items():
    p = next(p for p in doc.paragraphs if p.text.startswith(no + "\t"))
    kosongkan(p)
    runs(p, no, size=8)
    p.runs[-1].add_tab()
    runs(p, teks, size=8)

# =====================================================================================  BERSIH-BERSIH
for tag in ("w:commentRangeStart", "w:commentRangeEnd"):
    for el in list(body.iter(qn(tag))):
        el.getparent().remove(el)
for el in list(body.iter(qn("w:commentReference"))):
    r = el.getparent()
    r.getparent().remove(r)
cp = doc.part._comments_part
for c in list(cp.element):
    cp.element.remove(c)
for hl in list(body.iter(qn("w:highlight"))):
    par = next(a for a in hl.iterancestors() if a.tag == qn("w:p"))
    if "[ISIAN TIM" not in "".join(par.itertext()):
        hl.getparent().remove(hl)

doc.save(OUT)
SIDE.write_text(hashlib.sha256(OUT.read_bytes()).hexdigest())
print("ok", OUT.name)
