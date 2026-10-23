"""Angka kanon (§5 rencana) dan Tabel 9–11 dari eval/results.json (run 2) dan arsip run 1.

  python buat_angka_tabel.py   (dari folder Final/Paper)

Keluaran: ANGKA_KANON.md, Tabel_9_10_11.docx. Hanya data; tidak ada narasi.
"""
import copy
import json
import math
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent
EVAL = ROOT.parent / "Prototype" / "eval"
R2 = json.load(open(EVAL / "results.json", encoding="utf-8"))
R1 = json.load(open(EVAL / "results_run1_2026-10-12.json", encoding="utf-8"))
BASE = ROOT / "IRIS kata HQ ini lomba terakhir_ICC2026_Final.docx"


def wilson(k, n, z=1.959964):
    p = k / n
    c = (p + z * z / (2 * n)) / (1 + z * z / n)
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / (1 + z * z / n)
    return max(0.0, c - h), min(1.0, c + h)


def num(x, d=3):
    return f"{x:.{d}f}".replace(".", ",")


def kn(v):
    """'k/n [lo; hi]' format Indonesia; '—' bila tidak ada."""
    if not isinstance(v, dict) or "k" not in v:
        return "—" if v is None else str(v)
    return f"{v['k']}/{v['n']} [{num(v['ci95'][0])}; {num(v['ci95'][1])}]"


def get(r, path):
    for p in path.split("."):
        r = r.get(p) if isinstance(r, dict) else None
    return r


def cek_wilson(r):
    """Setiap ci95 di results.json harus sama dengan wilson(k, n) pada 4 desimal."""
    def walk(o):
        if isinstance(o, dict):
            if {"k", "n", "ci95"} <= o.keys():
                lo, hi = wilson(o["k"], o["n"])
                assert abs(lo - o["ci95"][0]) < 1e-4 and abs(hi - o["ci95"][1]) < 1e-4, o
                yield 1
            for v in o.values():
                yield from walk(v)
        elif isinstance(o, list):
            for v in o:
                yield from walk(v)
    return sum(walk(r))


# ---------- ANGKA_KANON.md ----------
KANON = [
    ("G1.D", "Dokumen uji D lulus G1", "IV.B, Tabel 10–11"),
    ("G1.B0", "B0 lulus G1", "Tabel 11"),
    ("G1.B1", "B1 lulus G1", "Tabel 11"),
    ("G1.D_ketat", "G1-ketat (penanda tahan dihitung belum bermakna)", "Tabel 10"),
    ("G1.veraPDF_lulus_D", "veraPDF penuh lulus (D)", "Tabel 11"),
    ("G3.D_fk_terbit", "Fk klaim terbit vs Tacuan", "Abstrak, IV.B, IV.C"),
    ("G3.D_terbit", "Grafik terbit otomatis (D)", "Tabel 11"),
    ("G3.D_R_beban_tinjauan", "R set uji", "IV.C RM3"),
    ("G3.injeksi_tertahan", "Injeksi tertahan", "IV.B, IV.C"),
    ("G3.injeksi_per_jenis.ubah_angka", "Injeksi: ubah angka", "IV.B"),
    ("G3.injeksi_per_jenis.tukar_label", "Injeksi: tukar label", "IV.B"),
    ("G3.injeksi_per_jenis.turunan_salah", "Injeksi: turunan salah", "IV.B"),
    ("G3.injeksi_per_jenis.tren_salah", "Injeksi: tren salah", "IV.B"),
    ("G3.B2_terbit", "B2 terbit", "Tabel 11"),
    ("G3.B2_angka_keliru_terbit", "Angka keliru B2 terbit (jumlah angka)", "Abstrak, IV.C"),
    ("G3.B2_grafik_dengan_angka_keliru", "Grafik B2 dengan angka keliru", "Abstrak, IV.C"),
    ("G3.D_angka_keliru_terbit", "Angka keliru D terbit", "Tabel 11"),
    ("G3.B2_fk_perkiraan", "Fk B2 (pemeriksa pola, perkiraan)", "Tabel 11"),
    ("G4.D", "G4 jawaban benar dari deskripsi D", "IV.B"),
    ("G4.B2", "G4 jawaban benar dari deskripsi B2", "IV.B"),
    ("G4.tabel", "G4 jawaban benar dari tabel (penyebut Q)", "IV.B"),
    ("G4.Q_D", "Rasio Q(D)", "IV.B"),
    ("G4.Q_B2", "Rasio Q(B2)", "IV.B"),
    ("G6_ringan.G1_D", "G1 dokumen nyata (D)", "IV.B, IV.D"),
    ("G6_ringan.G1_B0", "G1 dokumen nyata (B0)", "IV.B"),
    ("G6_ringan.G1_B1", "G1 dokumen nyata (B1)", "IV.B"),
    ("G6_ringan.veraPDF_lulus_D", "veraPDF penuh nyata (D)", "IV.B"),
    ("G6_ringan.R_beban_tinjauan", "R dokumen nyata", "IV.C, IV.D"),
]


def val(r, key):
    v = get(r, key)
    if isinstance(v, float):
        return num(v, 4).rstrip("0").rstrip(",")
    return kn(v)


def kanon_md():
    n_ci = cek_wilson(R2)
    a1, a2 = R1["ambang"], R2["ambang"]
    sama = all(a1[k] == a2[k] for k in ("theta_e", "epsilon", "rho", "locked_at"))
    pa = {p["id"]: p["status"] for p in R1["per_grafik"]}
    beda = [p["id"] for p in R2["per_grafik"] if pa.get(p["id"]) != p["status"]]
    rows = "\n".join(
        f"| `{k}` | {m} | {val(R1, k)} | {val(R2, k)} | {'sama' if get(R1, k) == get(R2, k) else ('baru di run 2' if get(R1, k) is None else '**BEDA**')} | {u} |"
        for k, m, u in KANON)
    per_dok = "\n".join(
        f"| {d['file']} | {d['halaman']} | {d['B0']['aturan_gagal']} | {d['B1']['aturan_gagal']} | {d['D']['aturan_gagal']} | {num(d['detik'], 1)} |"
        for d in R2["G1_per_dokumen"])
    g6 = "\n".join(
        f"| {d['file']} | {d['domain']} | {d['halaman']} | {d['gambar']} | {d['dekoratif']} | {d['figur_vektor_ditahan']} | {d['antrean']} | {d['terbit_otomatis']} | {d['B0']['aturan_gagal']} | {d['B1']['aturan_gagal']} | {d['D']['aturan_gagal']} | {num(d['detik'], 1)} |"
        for d in R2["G6_ringan"]["per_dokumen"])
    m = R2["manifest"]
    return f"""# Angka kanon paper final SIMAK (diisi otomatis)

Dibuat oleh `buat_angka_tabel.py` dari `Prototype/eval/results.json` (run 2, {R2['waktu']}) dan
`Prototype/eval/results_run1_2026-10-12.json` (run 1, {R1['waktu']}). Jalankan ulang skrip bila `results.json` berubah.
Format: k/n [batas bawah; batas atas] selang Wilson 95%, koma desimal.

## Pemeriksaan otomatis

- Selang Wilson: {n_ci} proporsi di run 2 dihitung ulang dengan pers. (8); semuanya cocok (4 desimal).
- Ambang run 1 = run 2: **{'ya' if sama else 'TIDAK'}** (θe = {num(a2['theta_e'], 2)}; ε = {num(a2['epsilon'], 2)}; ρ = {num(a2['rho'], 2)}; dikunci {a2['locked_at']}, {a2['calibrated_on']}).
- Status grafik berbeda run 1 vs run 2: **{len(beda)}/{len(R2['per_grafik'])}** {beda if beda else ''}.
- Manifes run 2: T1 `{m['model_t1']}`, T2 `{m['model_t2']}`, B2 `{m['model_b2']}`, QA `{m['model_qa']}`, suhu {m['temperature']},
  hash perintah T1 `{m['prompt_t1']}`, T2 `{m['prompt_t2']}`, B2 `{m['prompt_b2']}`, skema `{m['skema']}` (run 1: skema tidak tercatat).
- G6 ringan tidak diulang di run 2 (isi identik dengan run 1).

## Tabel §5

| Kunci (`results.json`) | Makna | Run 1 (12 Okt) | Run 2 (23 Okt) | Run 1 vs 2 | Dipakai di |
|---|---|---|---|---|---|
{rows}
| `ambang` | θe, ε, ρ | {num(a1['theta_e'], 2)}; {num(a1['epsilon'], 2)}; {num(a1['rho'], 2)} | {num(a2['theta_e'], 2)}; {num(a2['epsilon'], 2)}; {num(a2['rho'], 2)} | {'sama' if sama else '**BEDA**'} | III.G, Lampiran H |
| `eval/audit_b2.md` | Audit pemeriksa B2 | 13 pembulatan, 3 rerata keliru, 1 salah tandai | — | — | IV.B |
| `Semifinal/Audit/` | Audit awal | 0/28 PDF/UA-1; 29/1.111 Figure beralt | — | — | Abstrak, Bab I |

## Rincian per dokumen uji (run 2): aturan veraPDF gagal

| Dokumen | Halaman | B0 | B1 | D | Waktu proses (dtk) |
|---|---|---|---|---|---|
{per_dok}

## Rincian G6 ringan (6 PDF nyata)

| Dokumen | Domain | Hlm | Gambar | Dekoratif | Figur vektor ditahan | Antrean | Terbit | veraPDF B0 | B1 | D | Dtk |
|---|---|---|---|---|---|---|---|---|---|---|---|
{g6}
"""


# ---------- Tabel 9–11 (docx, gaya template: TNR 9 pt, keterangan di atas) ----------
TABEL9 = [
    ("ID", "Status", "Implementasi (berkas)", "Catatan realisasi", "Demo"),
    ("STR-01", "sebagian", "pdfa11y.autotag (OpenDataLoader); pipeline._blocks", "Cadangan tabel PyMuPDF find_tables, bukan Docling; tabel cadangan belum Table/TH di PDF", "ya"),
    ("CLS-01", "jalan", "pipeline._classify_decorative; kolom jenis VLM", "Heuristik ukuran, hash, pengulangan + klasifikasi VLM", "ya"),
    ("CHT-01", "jalan", "simak/vlm.py; chart.delta; chart.axis_check", "T1 dan T2 satu model (Gemma 4 26B-A4B), dua perintah", "ya"),
    ("CHT-02", "sebagian", "chart.derived; chart.build_description", "Statistik kode + templat level 1–3; tanpa VLM penyunting", "ya"),
    ("FRM-01", "belum", "gate.evaluate (aturan rumus)", "Tanpa pengenal kedua; semua rumus ditahan dengan draf LaTeX", "ya (penahanan)"),
    ("FRM-02", "belum", "—", "MathCAT belum diintegrasikan", "tidak"),
    ("GAT-01", "jalan", "gate.fk; gate.evaluate", "K = {cF, cδ, cA, cL}", "ya"),
    ("OUT-01", "jalan", "pdfa11y.patch; pdfa11y.verapdf; pdfa11y.g1_audit", "veraPDF 1.30.3; G1 dan G1-ketat; dekoratif belum Artifact", "ya"),
    ("OUT-02", "jalan", "simak/output.py; web/common.js", "Web Speech API id-ID bawaan; ElevenLabs opsional", "ya"),
    ("REV-01", "jalan", "pipeline.review; web/index.html", "Setujui/ubah/tolak/dekoratif; figur vektor belum tampil sebagai kartu", "ya"),
]


def tabel10():
    g1, g3, g4, g6 = R2["G1"], R2["G3"], R2["G4"], R2["G6_ringan"]
    k = R2["kriteria"]
    return [
        ("Gerbang", "Kriteria terkunci (Tabel 7)", "Hasil k/n [Wilson 95%]", "Status", "Data"),
        ("G1", "≥ 90% dokumen uji bebas kegagalan teks alternatif, bahasa, dan judul; lebih baik dari B1", f"D {kn(g1['D'])}; B1 {kn(g1['B1'])}", "lulus" if g1["lulus"] else "tidak lulus", "4 dokumen uji sintetis"),
        ("G1-ketat (tambahan)", "dilaporkan; bukan kriteria", f"D {kn(g1['D_ketat'])}", "dilaporkan", "4 dokumen uji sintetis"),
        ("G2", "τ ≥ 0,90", "—", "tidak dijalankan", "—"),
        ("G3 (Fk terbit)", "Fk ≥ 0,98 pada keluaran terbit", kn(g3["D_fk_terbit"]) + " klaim", "lulus" if g3["lulus"] else "tidak lulus", "60 grafik uji sintetis"),
        ("G3 (injeksi)", "≥ 0,90 deskripsi yang diubah tertahan", kn(g3["injeksi_tertahan"]), "lulus" if g3["lulus"] else "tidak lulus", "138 deskripsi diinjeksi"),
        ("G4", k["G4"], f"Q(D) = {num(g4['Q_D'], 2)} ({kn(g4['D'])}); Q(B2) = {num(g4['Q_B2'], 2)}", "lulus" if g4["lulus"] else "tidak lulus", "80 pertanyaan"),
        ("G5", "A ≥ 0,90", "—", "tidak dijalankan", "—"),
        ("G6 ringan", "G1 ≥ 80%; κ ≥ 0,60; R dilaporkan", f"G1 D {kn(g6['G1_D'])}; R {kn(g6['R_beban_tinjauan'])}; κ —", "sebagian (6 dari 60 dokumen; κ dan NVDA belum)", "6 PDF Common Crawl ac.id"),
        ("G7", "ditunda; disepakati dengan ULD", "—", "tidak dijalankan", "—"),
    ]


def tabel11():
    g1, g3, g4 = R2["G1"], R2["G3"], R2["G4"]
    d0 = R2["G1_per_dokumen"]
    aturan = lambda s: "; ".join(sorted({str(d[s]["aturan_gagal"]) for d in d0}))
    return [
        ("Ukuran", "B0", "B1", "B2", "D"),
        ("Lulus G1 (dokumen)", f"{g1['B0']['k']}/4", f"{g1['B1']['k']}/4", "—", f"{g1['D']['k']}/4"),
        ("Aturan veraPDF gagal per dokumen", aturan("B0"), aturan("B1"), "—", aturan("D")),
        ("Grafik terbit", "—", "—", f"{g3['B2_terbit']['k']}/60", f"{g3['D_terbit']['k']}/60"),
        ("Angka keliru terbit", "—", "—", f"{g3['B2_angka_keliru_terbit']} pada {g3['B2_grafik_dengan_angka_keliru']['k']}/60 grafik", str(g3["D_angka_keliru_terbit"])),
        ("Fk klaim terbit", "—", "—", f"{num(g3['B2_fk_perkiraan']['p'])} ({g3['B2_fk_perkiraan']['k']}/{g3['B2_fk_perkiraan']['n']}, perkiraan)", f"{num(g3['D_fk_terbit']['p'])} ({g3['D_fk_terbit']['k']}/{g3['D_fk_terbit']['n']})"),
        ("Rasio Q (G4)", "—", "—", num(g4["Q_B2"], 2), num(g4["Q_D"], 2)),
        ("Beban tinjauan R", "—", "—", "0/60 (tanpa gerbang)", f"{g3['D_R_beban_tinjauan']['k']}/60"),
    ]


def tulis_tabel(doc, judul, rows, catatan=None):
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(judul)
    r.font.name, r.font.size = "Times New Roman", Pt(9)
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t._tbl.tblPr.append(copy.deepcopy(BORDERS))  # border sama dengan tabel semifinal
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, txt in enumerate(row):
            p = t.cell(i, j).paragraphs[0]
            run = p.add_run(txt)
            run.font.name, run.font.size, run.bold = "Times New Roman", Pt(9), i == 0
    if catatan:
        p = doc.add_paragraph()
        run = p.add_run(catatan)
        run.font.name, run.font.size = "Times New Roman", Pt(8)
    doc.add_paragraph()


def tabel_docx(path):
    global BORDERS
    BORDERS = next(t._tbl.tblPr.find(qn("w:tblBorders")) for t in docx.Document(BASE).tables
                   if t._tbl.tblPr.find(qn("w:tblBorders")) is not None)
    doc = docx.Document()
    tulis_tabel(doc, "Tabel 9. Realisasi fitur Tabel 6 pada MVP", TABEL9)
    tulis_tabel(doc, "Tabel 10. Hasil gerbang terhadap kriteria lulus terkunci (run 2; run 1 identik)", tabel10())
    tulis_tabel(doc, "Tabel 11. SIMAK (D) dibandingkan dengan pembanding B0, B1, dan B2 pada set uji tertahan", tabel11(),
                "Fk B2 diukur dengan pemeriksa pola pada teks bebas (perkiraan; audit manual di eval/audit_b2.md).")
    doc.save(path)


if __name__ == "__main__":
    # self-check pers. (8) pada nilai yang dikenal: 4/4 → [0,510; 1,000]
    lo, hi = wilson(4, 4)
    assert round(lo, 4) == 0.5101 and round(hi, 4) == 1.0
    (ROOT / "ANGKA_KANON.md").write_text(kanon_md(), encoding="utf-8")
    tabel_docx(ROOT / "Tabel_9_10_11.docx")
    print("ok: ANGKA_KANON.md, Tabel_9_10_11.docx")
