"""Merakit naskah final dari docx semifinal: kerangka Bab IV–V, Tabel 9–11, Gambar 4, rujukan [55]–[56],
komentar revisi untuk tim, dan lampiran data H–P.

  python susun_final.py   (dari folder Final/Paper; jalankan buat_angka_tabel.py, lampiran/hitung_lampiran.py,
                           lampiran/buat_grafik.py, gambar/tangkap_mvp.py lebih dulu)

Hanya struktur, data, tabel, gambar, dan keterangan netral. Narasi (abstrak, analisis, pembahasan, kesimpulan,
revisi kalimat Bab I–III, Lampiran A) ditulis tim: Guidebook 7.2.4 butir 2. Penanda isian disorot kuning.
Menolak menimpa berkas final bila sudah disunting tangan sejak perakitan terakhir.
"""
import copy
import hashlib
import json
import sys
from importlib import metadata
from pathlib import Path

import docx
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

import buat_angka_tabel as bat

ROOT = Path(__file__).resolve().parent
PROTO = ROOT.parent / "Prototype"
SRC = ROOT.parent / "Semifinal/Draft/Fix/IRIS kata HQ ini lomba terakhir_ICC2026_Semifinal.docx"
OUT = ROOT / "IRIS kata HQ ini lomba terakhir_ICC2026_Final.docx"
SIDE = ROOT / ".rakitan.sha256"
R2, R1 = bat.R2, bat.R1
L = json.load(open(ROOT / "lampiran/lampiran_data.json", encoding="utf-8"))
GL, GB = ROOT / "lampiran/gambar", ROOT / "gambar"
num, kn = bat.num, bat.kn
AUTHOR, INIT = "Claude (asisten AI)", "AI"
FULL = 10368  # lebar teks satu kolom potret (dxa)

doc = docx.Document(SRC)
body = doc.element.body
B = list(body)
SECT2 = copy.deepcopy(B[12].find(".//" + qn("w:sectPr")))   # akhir bagian dua kolom
SECT1 = copy.deepcopy(B[14].find(".//" + qn("w:sectPr")))   # akhir bagian satu kolom (kontinu)
TBLPR = copy.deepcopy(doc.tables[8]._tbl.tblPr)             # gaya Tabel 6 semifinal
paras = {id(p._p): p for p in doc.paragraphs}


def find_p(sub):
    """Paragraf pertama (badan atau sel tabel) yang memuat sub; spasi tak putus dinormalkan."""
    cells = (p for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs)
    for p in (*doc.paragraphs, *cells):
        if sub in p.text.replace("\xa0", " "):
            return p
    raise SystemExit(f"teks tidak ditemukan: {sub!r}")


# ------------------------------------------------------------------ pembangun elemen
def _run(p, text, size=None, bold=False, italic=False, hl=False):
    r = p.add_run(text)
    if size:
        r.font.size = Pt(size)
    r.bold, r.italic = bold or None, italic or None
    if hl:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return r


def _pPr(p):
    return p._p.get_or_add_pPr()


def para(text="", *, indent=True, jc="both", size=None, italic=False, hl=False, keep=False, after=None):
    p = doc.add_paragraph()
    pp = _pPr(p)
    if indent:
        ind = OxmlElement("w:ind"); ind.set(qn("w:firstLine"), "142"); pp.append(ind)
    if keep:
        pp.append(OxmlElement("w:keepNext"))
    if after is not None:
        sp = OxmlElement("w:spacing"); sp.set(qn("w:after"), str(after)); pp.append(sp)
    j = OxmlElement("w:jc"); j.set(qn("w:val"), jc); pp.append(j)
    if text:
        _run(p, text, size=size, italic=italic, hl=hl)
    return p


def isian(teks):
    """Penanda isian untuk tim (kuning, miring)."""
    return para(f"[ISIAN TIM — {teks} Hapus penanda ini setelah diisi.]", italic=True, hl=True)


def h1(text, numbered=True, page_break=False):
    p = doc.add_paragraph(style="Heading 1")
    pp = _pPr(p)
    if numbered:
        n = OxmlElement("w:numPr")
        for tag, v in (("w:ilvl", "0"), ("w:numId", "2")):
            e = OxmlElement(tag); e.set(qn("w:val"), v); n.append(e)
        pp.append(n)
        _run(p, text[0]); _run(p, text[1:], size=8)   # kapital kecil seperti judul bab semifinal
    else:
        if page_break:
            pp.append(OxmlElement("w:pageBreakBefore"))
        ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "0"); ind.set(qn("w:firstLine"), "0"); pp.append(ind)
        _run(p, text)
    return p


def h2(text):
    p = doc.add_paragraph(style="Heading 2")
    n = OxmlElement("w:numPr")
    for tag, v in (("w:ilvl", "1"), ("w:numId", "2")):
        e = OxmlElement(tag); e.set(qn("w:val"), v); n.append(e)
    _pPr(p).append(n)
    _run(p, text)
    return p


def _cell_p(cell, text, *, bold=False, jc="left", size=9, keep=False, italic=False, hl=False):
    p = cell.paragraphs[0]
    pp = _pPr(p)
    if keep:
        pp.append(OxmlElement("w:keepNext"))
    sp = OxmlElement("w:spacing"); sp.set(qn("w:line"), "228"); sp.set(qn("w:lineRule"), "auto"); pp.append(sp)
    j = OxmlElement("w:jc"); j.set(qn("w:val"), jc); pp.append(j)
    _run(p, text, size=size, bold=bold, italic=italic, hl=hl)


def _row_flags(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))
    if header:
        trPr.append(OxmlElement("w:tblHeader"))


def tabel(caption, header, rows, widths, align=None, note=None, hl_col=None):
    """Tabel bergaya semifinal: baris keterangan (gabung, tanpa garis) + baris judul berulang."""
    n = len(header)
    assert len(widths) == n and all(len(r) == n for r in rows), caption
    scale = FULL / sum(widths) if sum(widths) > FULL else 1
    widths = [int(w * scale) for w in widths]
    t = doc.add_table(rows=len(rows) + 2, cols=n)
    tbl = t._tbl
    tbl.remove(tbl.tblPr)
    pr = copy.deepcopy(TBLPR)
    pr.find(qn("w:tblW")).set(qn("w:w"), str(sum(widths)))
    tbl.insert(0, pr)
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    for r in t.rows:
        for c, w in zip(r.cells, widths):
            c.width = docx.shared.Twips(w)
    cap = t.cell(0, 0).merge(t.cell(0, n - 1))
    tcPr = cap._tc.get_or_add_tcPr()
    bd = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "nil"); bd.append(e)
    tcPr.append(bd)
    p = cap.paragraphs[0]
    pp = _pPr(p)
    pp.append(OxmlElement("w:keepNext"))
    sp = OxmlElement("w:spacing"); sp.set(qn("w:before"), "100"); sp.set(qn("w:after"), "60"); pp.append(sp)
    j = OxmlElement("w:jc"); j.set(qn("w:val"), "center"); pp.append(j)
    _run(p, caption, size=9)
    _row_flags(t.rows[0], header=True)
    _row_flags(t.rows[1], header=True)
    for c, txt in zip(t.rows[1].cells, header):
        _cell_p(c, txt, bold=True, jc="center", keep=True)
        c.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER
    align = align or ["left"] * n
    for r, data in zip(t.rows[2:], rows):
        _row_flags(r)
        for j_, (c, txt) in enumerate(zip(r.cells, data)):
            _cell_p(c, str(txt), jc=align[j_], hl=hl_col is not None and j_ == hl_col)
    # paragraf pemisah wajib: dua tabel yang bersebelahan langsung digabung Word menjadi satu
    return [tbl, para(note or "", indent=False, jc="left", size=8, after=160)._p]


def gambar(path, caption, width_in):
    """Gambar sebaris dalam tabel satu kolom tanpa garis, keterangan di bawah (seperti semifinal)."""
    t = doc.add_table(rows=2, cols=1)
    p = t.cell(0, 0).paragraphs[0]
    pp = _pPr(p)
    pp.append(OxmlElement("w:keepNext"))
    j = OxmlElement("w:jc"); j.set(qn("w:val"), "center"); pp.append(j)
    pic = p.add_run().add_picture(str(path), width=Inches(width_in))
    pic._inline.docPr.set("descr", caption)
    _row_flags(t.rows[0]); _row_flags(t.rows[1])
    _cell_p(t.cell(1, 0), caption, jc="center")
    _pPr(t.cell(1, 0).paragraphs[0]).find(qn("w:spacing")).set(qn("w:after"), "120")
    return [t._tbl, para("", indent=False, after=0)._p]


def sisip(anchor, items):
    """Sisipkan elemen (paragraf/tabel/list) tepat sebelum anchor, berurutan."""
    for it in items:
        if isinstance(it, list):
            sisip(anchor, it)
            continue
        el = it._p if isinstance(it, Paragraph) else it
        anchor.addprevious(el)


def lebar(prev_para, objs):
    """Objek selebar halaman di tengah aliran dua kolom: tutup bagian dua kolom, objek, tutup satu kolom."""
    _pPr(prev_para).append(copy.deepcopy(SECT2))
    tutup = doc.add_paragraph()
    _pPr(tutup).append(copy.deepcopy(SECT1))
    return [prev_para, *objs, tutup]


def komentar(target, teks):
    p = target if isinstance(target, Paragraph) else find_p(target)
    runs = [r for r in p.runs if r.text.strip()] or p.runs
    doc.add_comment(runs[:1], text=teks, author=AUTHOR, initials=INIT)


def idn_p(v, d=3):
    return num(v, d) if v is not None else "—"


# ------------------------------------------------------------------ Bab IV–V (kerangka) sebelum REFERENSI
ref = find_p("REFERENSI")
g4 = gambar(GB / "Gambar4_montase.png",
            "Gambar 4. Antarmuka MVP SIMAK: (a) Remediasi & Tinjauan dengan validasi veraPDF B0, B1, D dan antrean "
            "tinjauan ULD; (b) Demo Gerbang pada grafik syn043 setelah injeksi tukar label–nilai; (c) halaman Evaluasi; "
            "(d) HTML keluaran dengan penanda status dan bilah mode dengar.", 7.0)
g5 = gambar(GB / "Gambar5_alur_gerbang_syn043.png",
            "Gambar 5. Alur gerbang SIMAK pada grafik uji syn043: deskripsi D terbit setelah empat pemeriksaan lulus; "
            "injeksi tukar label–nilai pada deskripsi yang sama ditahan; deskripsi B2 dari model yang sama terbit "
            "dengan satu angka rerata yang tidak setia.", 7.0)
g6 = gambar(GB / "Gambar6_status_D_vs_B2.png",
            "Gambar 6. Status keluaran 60 grafik set uji tertahan pada B2 (VLM langsung) dan D (SIMAK).", 7.0)
g7 = gambar(GB / "Gambar7_gerbang_vs_kriteria.png",
            "Gambar 7. Proporsi dan selang Wilson 95% per gerbang untuk D dan pembanding; garis putus menandai "
            "kriteria lulus terkunci (Tabel 7).", 7.0)
t9 = tabel("Tabel 9. Realisasi fitur Tabel 6 pada MVP", bat.TABEL9[0], [list(r) for r in bat.TABEL9[1:]],
           [800, 900, 3100, 4168, 1400])
t10 = tabel("Tabel 10. Hasil gerbang terhadap kriteria lulus terkunci (run 2; run 1 identik)", bat.tabel10()[0],
            [list(r) for r in bat.tabel10()[1:]], [1500, 2900, 2900, 1700, 1368])
t11 = tabel("Tabel 11. SIMAK (D) dibandingkan dengan pembanding B0, B1, dan B2 pada set uji tertahan",
            bat.tabel11()[0], [list(r) for r in bat.tabel11()[1:]], [2900, 1300, 1300, 2600, 2268],
            align=["left", "center", "center", "center", "center"],
            note="Fk B2 diukur dengan pemeriksa pola pada teks bebas (perkiraan; audit manual di Lampiran L). "
                 "Rincian per grafik di Lampiran J; selang Wilson seluruh ukuran di Lampiran I.")

ia = isian("IV.A Hasil Implementasi: apa yang dibangun, dirujuk ke Gambar 4, Gambar 5, dan Tabel 9 sebelum objek "
           "muncul. Lihat BRIEF_SISA_PENULISAN.md §IV.A.")
ib = isian("IV.B Pengujian: hasil metrik Bab III dengan kriteria terkunci; rujuk Gambar 6, Gambar 7, Tabel 10, dan "
           "Tabel 11; angka dari ANGKA_KANON.md. Lihat BRIEF_SISA_PENULISAN.md §IV.B.")
sisip(ref._p, [
    h1("HASIL DAN PEMBAHASAN"),
    h2("Hasil Implementasi"),
    lebar(ia, g4 + g5 + t9),
    h2("Pengujian"),
    lebar(ib, g6 + g7 + t10 + t11),
    h2("Pembahasan"),
    isian("IV.C Pembahasan: jawab RM1, RM2, RM3 dengan bukti dari IV.B; bandingkan dengan pustaka Bab II."),
    h2("Keterbatasan"),
    isian("IV.D Keterbatasan: bagian yang belum berjalan atau belum diuji (lihat status 'belum' di Tabel 9 dan "
          "'tidak dijalankan' di Tabel 10)."),
    h1("PENUTUP"),
    h2("Kesimpulan"),
    isian("V.A Kesimpulan: jawaban singkat atas setiap rumusan masalah."),
    h2("Rencana Pengembangan"),
    isian("V.B Rencana Pengembangan: fitur atau langkah berikutnya setelah MVP."),
    h2("Dampak"),
    isian("V.C Dampak: manfaat bagi target pengguna bila solusi dipakai."),
])

# ------------------------------------------------------------------ rujukan [55] dan [56]
r54 = next(p for p in doc.paragraphs if p.text.startswith("[54]\t"))  # entri pustaka, bukan kutipan di teks
REF_BARU = [
    ("[55]", "T. Allison, M. Milano, R. Stonebraker, dan P. Wyatt, “CC-MAIN-2021-31-PDF-UNTRUNCATED,” korpus SafeDocs, "
             "Digital Corpora. Diakses: 23 Okt. 2026. Tersedia: corp.digitalcorpora.org/corpora/files/CC-MAIN-2021-31-PDF-UNTRUNCATED/"),
    ("[56]", "A. Rossberg, Ed., “WebAssembly core specification,” W3C Recommendation, 5 Des. 2019. "
             "Tersedia: www.w3.org/TR/wasm-core-1/"),
]
anchor = r54._p.getnext()
for no, teks in REF_BARU:
    p = doc.add_paragraph()
    ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "270"); ind.set(qn("w:hanging"), "295"); _pPr(p).append(ind)
    _run(p, no, size=8, hl=True); p.add_run().add_tab(); _run(p, teks, size=8, hl=True)
    anchor.addprevious(p._p)
    komentar(p, "Rujukan ini ditemukan lewat penelusuran awal berbantuan AI. Wajib diverifikasi tim dari sumber "
                "aslinya sebelum dipakai (7.2.4 butir 2); hapus sorotan setelah diverifikasi.")

# ------------------------------------------------------------------ Tabel 6: kolom status diperbarui (data)
T6 = doc.tables[8]
assert T6.rows[0].cells[0].text.startswith("Tabel 6")
status6 = {r[0]: r[1] for r in bat.TABEL9[1:]}
label6 = {"jalan": "Berjalan pada MVP", "sebagian": "Berjalan sebagian; lihat Tabel 9", "belum": "Belum diimplementasikan; ditahan (fail-closed)"}
for row in T6.rows[2:]:
    fid = row.cells[0].text.strip()
    if fid in status6:
        c = row.cells[4]
        for r in c.paragraphs[0].runs:
            r.text = ""
        _run(c.paragraphs[0], label6[status6[fid]], size=9, hl=True)
komentar(T6.rows[0].cells[0].paragraphs[0],
         "Kolom Status diganti dari status 30 Sep ke status MVP (sumber: kode dan Tabel 9). Kolom Model/metode "
         "belum diubah: CHT-02 (VLM penyunting) dan STR-01 (Docling) berbeda dengan kode; revisi kalimat oleh tim.")

# ------------------------------------------------------------------ komentar revisi minor (tim menulis kalimatnya)
REVISI = [
    ("Abstrak", "Abstrak final harus memuat hasil utama; pernyataan bahwa makalah belum melaporkan kinerja harus "
                "dihapus (rencana §4 Abstrak). Ditulis tim, ≤ 250 kata, angka dari ANGKA_KANON.md."),
    ("Tabel 3. Manfaat", "Kolom Status masih per 30 Sep; perbarui ke hasil nyata atau 'tidak dijalankan' (rencana §4 I.D)."),
    ("SIMAK dibangun sebagai aplikasi web tiga lapis", "Sebut model nyata: Gemma 4 26B-A4B lewat Gemini API, suhu 0, "
                "cache per hash; T1 dan T2 satu model dengan dua perintah; mode dengar Web Speech API id-ID bawaan, "
                "ElevenLabs opsional (rencana §4 III.B)."),
    ("Modul grafik.", "Kode tidak memakai VLM penyunting kalimat: deskripsi murni templat (A18). Revisi oleh tim."),
    ("Modul rumus.", "Kode belum memakai pengenal kedua dan MathCAT; semua rumus ditahan (fail-closed) (A16)."),
    ("Sistem D dibandingkan dengan tiga pembanding", "Kriteria Tabel 7 tidak diubah. G1-ketat disebut sebagai "
                "pelaporan tambahan, bukan kriteria baru (A5)."),
    ("Rencana kerja (Lampiran F.13)", "Jadwal 1–23 Okt sudah lewat; ganti dengan satu kalimat realisasi atau hapus "
                "(rencana §4 III.I)."),
    ("Tabel 5. Dataset", "Tetap sebagai rancangan; rujuk tabel rencana vs realisasi di Lampiran O (B5)."),
    ("Tim menggunakan Claude", "Lampiran A wajib ditulis ulang tim sesuai penggunaan AI sebenarnya pada versi final, "
                "termasuk bantuan Claude untuk skrip evaluasi, angka kanon, Tabel 9–11, Gambar 4, lampiran data H–P, "
                "penelusuran rujukan, dan komentar revisi ini (7.2.4 butir 3)."),
    ("Perangkat: laptop tanpa GPU", "Uji kelayakan awal memakai veraPDF 1.30.2; MVP final memakai 1.30.3 (Lampiran H)."),
    ("Antarmuka: HTML, CSS, dan JavaScript", "Perbarui ke kenyataan MVP: model Gemma 4 26B-A4B, Web Speech API, "
                "titik akhir di Lampiran P; MathCAT belum dipakai (A4)."),
]
for sub, teks in REVISI:
    komentar(sub, teks)
komentar("Gerbang verifikasi. Setiap elemen e memiliki himpunan pemeriksaan K(e)", "Pers. (7): himpunan pemeriksaan grafik di kode adalah {cF, cδ, cA, cL}; tambahkan definisi cL dan "
             "asal-usulnya dari kalibrasi set pengembangan (Lampiran M; simak/gate.py). Pertahankan kalimat bahwa "
             "Fk terhadap T bernilai 1 menurut definisi (A10).")
pdl = [p for p in doc.paragraphs if "Docling" in p.text and not p.text.startswith("[")]
if pdl:
    komentar(pdl[0], "Cadangan tabel yang benar-benar dipakai kode adalah PyMuPDF find_tables, bukan Docling "
                     "(simak/pipeline.py _blocks).")
ps = doc.paragraphs
pc = next(ps[i + 1] for i, p in enumerate(ps) if p.style.name == "Heading 2" and p.text.strip() == "Alur Penelitian")
if pc is not None:
    komentar(pc, "Janji dua run sudah dipenuhi: run 1 (12 Okt) dan run 2 (23 Okt), 0/60 status grafik berbeda (Lampiran H).")

# ------------------------------------------------------------------ Lampiran B: penanda isian Turnitin
lb = find_p("Lampiran B. Hasil Pengecekan Similarity")
lb._p.getnext().addnext(para("[ISIAN TIM — gambar di atas adalah laporan Turnitin semifinal. Ganti dengan halaman "
                             "ringkasan laporan Turnitin naskah final (< 15%, di luar kutipan dan daftar pustaka); simpan "
                             "laporannya di Paper/turnitin/. Hapus penanda ini.]", italic=True, hl=True)._p)

# ------------------------------------------------------------------ lampiran baru H–P (bagian potret baru)
final_sect = body.find(qn("w:sectPr"))
tutup_lanskap = doc.add_paragraph()
_pPr(tutup_lanskap).append(copy.deepcopy(final_sect))       # Lampiran F–G tetap lanskap
pg = final_sect.find(qn("w:pgSz"))
pg.set(qn("w:w"), "12240"); pg.set(qn("w:h"), "15840")
pg.attrib.pop(qn("w:orient"), None)
END = final_sect
akhir = []


def lampiran(judul, sumber, *isi, baru_hal=True):
    akhir.extend([h1(judul, numbered=False, page_break=baru_hal),
                  para(sumber, indent=False, jc="left", size=8, italic=True, after=120), *isi])


m1, m2 = R1["manifest"], R2["manifest"]
ver = {p: metadata.version(p) for p in ("pymupdf", "opendataloader-pdf", "pikepdf", "fastapi", "google-genai", "matplotlib")}

# H. Reproduksi
lampiran(
    "Lampiran H. Reproduksi dan Manifes Evaluasi",
    "Sumber: github.com/Faiz7557/IRIS-Pusaka, tag v1.0-final; eval/results_run1_2026-10-12.json (run 1) dan "
    "eval/results.json (run 2).",
    tabel("Tabel H.1. Manifes kedua run evaluasi", ["Butir", "Run 1", "Run 2"], [
        ["Waktu selesai", R1["waktu"], R2["waktu"]],
        ["Model T1 / T2", f"{m1['model_t1']} / {m1['model_t2']}", f"{m2['model_t1']} / {m2['model_t2']}"],
        ["Model B2 / penjawab G4", f"{m1['model_b2']} / {m1['model_qa']}", f"{m2['model_b2']} / {m2['model_qa']}"],
        ["Suhu", str(m1["temperature"]), str(m2["temperature"])],
        ["Hash perintah T1 / T2 / B2", f"{m1['prompt_t1']} / {m1['prompt_t2']} / {m1['prompt_b2']}",
         f"{m2['prompt_t1']} / {m2['prompt_t2']} / {m2['prompt_b2']}"],
        ["Hash skema ekstraksi", m1.get("skema", "tidak tercatat"), m2["skema"]],
        ["Data uji", f"{R1['data']['grafik_uji']} grafik, {R1['data']['dokumen_uji']} dokumen; hash {R1['data']['hash_charts_json']}; seed {R1['data']['seed']}",
         f"{R2['data']['grafik_uji']} grafik, {R2['data']['dokumen_uji']} dokumen; hash {R2['data']['hash_charts_json']}; seed {R2['data']['seed']}"],
        ["Ambang θe; ε; ρ", f"{num(R1['ambang']['theta_e'], 2)}; {num(R1['ambang']['epsilon'], 2)}; {num(R1['ambang']['rho'], 2)}",
         f"{num(R2['ambang']['theta_e'], 2)}; {num(R2['ambang']['epsilon'], 2)}; {num(R2['ambang']['rho'], 2)}"],
        ["Ambang dikunci", R1["ambang"]["locked_at"], R2["ambang"]["locked_at"]],
    ], [2600, 3884, 3884]),
    tabel("Tabel H.2. Perintah reproduksi (dari folder Prototype)", ["Langkah", "Perintah", "Keluaran"], [
        ["1. Pasang dependensi", "pip install -r requirements.txt", "Paket Python inti"],
        ["2. Bangkitkan set sintetis", "python data/make_synth.py 90", "90 grafik, 6 dokumen PDF"],
        ["3. Kalibrasi (set dev saja)", "python -m eval.evaluate kalibrasi", "simak/thresholds.json (terkunci)"],
        ["4. Evaluasi set uji tertahan", "python -m eval.evaluate uji", "eval/results.json"],
        ["5. Selisih dua run", "python -m eval.banding_run eval/results_run1_2026-10-12.json eval/results.json", "Selisih per ukuran dan per grafik"],
        ["6. Audit pemeriksa B2", "python eval/audit_b2.py", "Daftar angka B2 yang ditandai"],
        ["7. Aplikasi demo", "python -m uvicorn app:app --port 8000", "http://localhost:8000"],
        ["8. Data lampiran ini", "python lampiran/hitung_lampiran.py (folder Final/Paper)", "lampiran_data.json; status 60 grafik dicocokkan dengan run 2"],
    ], [2300, 5000, 3068]),
    tabel("Tabel H.3. Ukuran utama pada run 1 dan run 2", ["Ukuran", "Run 1", "Run 2", "Sama"], [
        [m, bat.val(R1, k), bat.val(R2, k), "ya" if bat.get(R1, k) == bat.get(R2, k) else ("baru" if bat.get(R1, k) is None else "TIDAK")]
        for k, m, _ in bat.KANON], [4000, 2700, 2700, 968], align=["left", "center", "center", "center"],
        note=f"Status grafik berbeda antara run 1 dan run 2: "
             f"{sum(1 for p in R2['per_grafik'] if {q['id']: q['status'] for q in R1['per_grafik']}[p['id']] != p['status'])}/60."),
    tabel("Tabel H.4. Lingkungan perangkat lunak", ["Komponen", "Versi"], [
        ["Python", sys.version.split()[0]], ["Java (veraPDF)", "21.0.8"], ["veraPDF", "1.30.3 (profil PDF/UA-1)"],
        ["OpenDataLoader PDF", ver["opendataloader-pdf"]], ["PyMuPDF", ver["pymupdf"]], ["pikepdf", ver["pikepdf"]],
        ["FastAPI", ver["fastapi"]], ["google-genai", ver["google-genai"]], ["matplotlib (grafik lampiran)", ver["matplotlib"]],
    ], [5184, 5184]),
)

# I. Ringkasan selang Wilson
semua = []
for k, m, _ in bat.KANON:
    v = bat.get(R2, k)
    if isinstance(v, dict) and "k" in v:
        semua.append([m, f"{v['k']}/{v['n']}", idn_p(v["p"]), f"[{num(v['ci95'][0])}; {num(v['ci95'][1])}]"])
lampiran(
    "Lampiran I. Seluruh Ukuran dengan Selang Wilson 95%",
    "Sumber: eval/results.json (run 2). Selang dihitung dengan pers. (8) dan dicek ulang oleh buat_angka_tabel.py.",
    gambar(GL / "L_selang_wilson.png", "Gambar I.1. Proporsi dan selang Wilson 95% per ukuran untuk D dan pembanding; "
           "garis putus vertikal menandai kriteria lulus terkunci.", 7.0),
    tabel("Tabel I.1. Proporsi dan selang Wilson 95% seluruh ukuran (run 2)", ["Ukuran", "k/n", "p", "Selang Wilson 95%"],
          semua, [5000, 1500, 1400, 2468], align=["left", "center", "center", "center"]),
)

# J. Per grafik
CK = ("cF", "cδ", "cA", "cL")
JN = {"batang": "batang", "batang_horizontal": "batang horiz.", "garis": "garis", "area": "area", "titik": "titik"}
rows_j = []
for g in L["grafik"]:
    ck = " ".join(("✓" if g["checks"].get(c) else "✗") for c in CK)
    rows_j.append([g["id"], g["judul"], JN[g["jenis"]], "ya" if g["label_nilai"] else "tidak", str(g["n_kategori"]),
                   num(g["delta"], 4) if g["delta"] is not None else "—", ck,
                   f"{g.get('klaim_setia', 0)}/{g['n_klaim']}", "terbit" if g["status"] == "terbit_otomatis" else "ditahan",
                   f"{g['b2_angka']} ({g['b2_keliru']})"])
lampiran(
    "Lampiran J. Hasil per Grafik pada Set Uji Tertahan",
    "Sumber: lampiran_data.json, dihitung ulang luring dari cache VLM; status setiap grafik identik dengan eval/results.json (run 2).",
    gambar(GL / "L_status_grafik.png", "Gambar J.1. Status keluaran SIMAK pada 60 grafik uji menurut (a) jenis grafik "
           "dan (b) keberadaan label nilai.", 7.0),
    tabel("Tabel J.1. Hasil per grafik: pemeriksaan gerbang, kesetiaan klaim, status, dan angka B2",
          ["ID", "Judul grafik", "Jenis", "Label nilai", "Kat.", "δ", "cF cδ cA cL", "Klaim setia", "Status", "B2 angka (keliru)"],
          rows_j, [780, 2720, 950, 650, 450, 650, 1150, 750, 750, 900],
          align=["left", "left", "left", "center", "center", "center", "center", "center", "center", "center"],
          note="Kolom cF cδ cA cL: ✓ lulus, ✗ gagal (pers. (7)). Klaim setia: klaim deskripsi D yang cocok dengan tabel "
               "acuan pada toleransi ε. B2 angka: angka pada deskripsi B2 yang diperiksa (dalam kurung: ditandai tidak setia)."),
)

# K. Injeksi
NJ = {"ubah_angka": "Ubah angka", "tukar_label": "Tukar label", "turunan_salah": "Turunan salah", "tren_salah": "Tren salah"}
per = {}
for i in L["injeksi"]:
    per.setdefault(i["id"], {})[i["jenis"]] = i
rows_k = [[cid, *[("tertahan" if v[k]["tertahan"] else "LOLOS") if k in v else "—" for k in NJ],
           num(next(iter(v.values()))["eps_s"], 2)] for cid, v in per.items()]
contoh = []
for k in NJ:
    ex = next(i for i in L["injeksi"] if i["jenis"] == k and i["tertahan"])
    contoh.append([NJ[k], ex["id"], ex["alasan"][0] if ex["alasan"] else "—"])
lolos = [i for i in L["injeksi"] if not i["tertahan"]]
lampiran(
    "Lampiran K. Rincian Uji Injeksi Kesalahan",
    f"Sumber: lampiran_data.json; injeksi diterapkan pada {len(per)} grafik terbit dengan seed 20260930 dan ε = 0,02 "
    f"(pers. (5)); jumlah tertahan identik dengan eval/results.json.",
    gambar(GL / "L_injeksi.png", "Gambar K.1. Proporsi deskripsi terinjeksi yang tertahan per jenis injeksi dengan selang Wilson 95%.", 7.0),
    tabel("Tabel K.1. Hasil injeksi per grafik terbit dan per jenis", ["Grafik", *NJ.values(), "ε·s"], rows_k,
          [1400, 1900, 1900, 1900, 1900, 1368], align=["left", "center", "center", "center", "center", "center"],
          note="— : jenis injeksi tidak berlaku (mis. grafik tanpa urutan untuk tren). ε·s: toleransi absolut "
               "(ε dikali rentang nilai grafik)." + (f" Kasus lolos: {', '.join(i['id'] + ' (' + NJ[i['jenis']] + ')' for i in lolos)}." if lolos else "")),
    tabel("Tabel K.2. Contoh alasan penahanan yang dicatat gerbang untuk setiap jenis injeksi", ["Jenis", "Grafik", "Alasan tercatat"],
          contoh, [1800, 1200, 7368]),
)

# L. Audit B2
def kat_b2(a):
    if a["id"] == "syn079":
        return "Salah tandai pemeriksa"
    if a["id"] in ("syn043", "syn083") or (a["id"] == "syn068" and "rata" in a["kutipan"]):
        return "Rerata dihitung model keliru"
    return "Pembulatan/perkiraan di luar ε"


VD = {"dikarang": "tidak cocok dengan acuan", "label_tertukar": "label tertukar"}
lampiran(
    "Lampiran L. Audit Manual Pemeriksa Angka B2",
    "Sumber: eval/audit_b2.md dan lampiran_data.json (17 angka B2 yang ditandai pemeriksa pola pada 60 grafik uji).",
    gambar(GL / "L_audit_b2.png", "Gambar L.1. Angka B2 yang ditandai tidak setia per grafik menurut kategori audit manual.", 7.0),
    tabel("Tabel L.1. Daftar angka B2 yang ditandai dan kategori audit manualnya",
          ["No", "Grafik", "Label nilai", "Putusan pemeriksa", "Kutipan deskripsi B2", "Kategori audit"],
          [[str(i + 1), a["id"], "ya" if a["label_nilai"] else "tidak", VD.get(a["verdict"], a["verdict"]),
            "…" + a["kutipan"] + "…", kat_b2(a)] for i, a in enumerate(L["audit_b2"])],
          [450, 800, 700, 1400, 4818, 2200], align=["center", "left", "center", "left", "left", "left"],
          note="Kategori dari audit manual tim (eval/audit_b2.md): 13 pembulatan/perkiraan di luar ε, 3 rerata yang "
               "dihitung model keliru, 1 salah tandai pemeriksa."),
)

# M. Kalibrasi
v1g, v2g = L["kalibrasi"]["v1"]["grid"], L["kalibrasi"]["v2"]["grid"]
lampiran(
    "Lampiran M. Kalibrasi Ambang pada Set Pengembangan",
    "Sumber: eval/kalibrasi_v1_sebelum_perbaikan.json, eval/kalibrasi.json, dan simak/thresholds.json "
    "(30 grafik dev; set uji belum dibuka saat kalibrasi).",
    gambar(GL / "L_kalibrasi.png", "Gambar M.1. Jumlah grafik dev yang terbit dan rerata Fk terhadap acuan menurut "
           "ambang θe, versi 1 dan versi 2.", 7.0),
    tabel("Tabel M.1. Kisi kalibrasi θe pada 30 grafik dev", ["θe", "Terbit v1", "Rerata Fk v1", "Terbit v2", "Rerata Fk v2"],
          [[num(a["theta_e"], 3), str(a["terbit"]), num(a["fk_rerata"], 4), str(b["terbit"]), num(b["fk_rerata"], 4)]
           for a, b in zip(v1g, v2g)], [1800, 2100, 2200, 2100, 2168], align=["center"] * 5),
    tabel("Tabel M.2. Aturan dan riwayat ambang", ["Versi", "Dikunci", "θe", "Aturan pemilihan", "Riwayat"], [
        ["v1", L["kalibrasi"]["v1"]["terkunci"]["locked_at"], num(L["kalibrasi"]["v1"]["terkunci"]["theta_e"], 2),
         L["kalibrasi"]["v1"]["terkunci"]["aturan"], "—"],
        ["v2", L["kalibrasi"]["v2"]["terkunci"]["locked_at"], num(L["kalibrasi"]["v2"]["terkunci"]["theta_e"], 2),
         L["kalibrasi"]["v2"]["terkunci"]["aturan"], L["kalibrasi"]["v2"]["terkunci"]["riwayat"]],
    ], [600, 1500, 600, 3800, 3868]),
    tabel("Tabel M.3. Hasil per grafik dev pada versi 2", ["ID", "δ", "cA (sumbu)", "Label nilai", "Fk vs acuan"],
          [[g["id"], num(g["delta"], 4), "✓" if g["axis"] else "✗", "ya" if g["label"] else "tidak", num(g["fk"], 3)]
           for g in L["kalibrasi"]["v2"]["per_grafik"]], [2000, 2000, 2000, 2000, 2368], align=["left"] + ["center"] * 4),
)

# N. G6 ringan
G6 = R2["G6_ringan"]
cc = {s["file"]: s for s in json.load(open(PROTO / "data/cc/cc.json", encoding="utf-8"))["sampel"]}
lampiran(
    "Lampiran N. Rincian G6 Ringan pada PDF Nyata",
    f"Sumber: eval/results.json (G6_ringan) dan data/cc/cc.json; {G6['sumber']}. Berkas PDF tidak disebarkan.",
    gambar(GL / "L_g6_verapdf.png", "Gambar N.1. Jumlah aturan PDF/UA-1 yang gagal menurut veraPDF pada enam PDF nyata untuk B0, B1, dan D.", 7.0),
    tabel("Tabel N.1. Rincian per dokumen G6 ringan",
          ["Dokumen", "Domain", "Hlm", "Bertag asli", "Gambar", "Dekoratif", "Figur vektor ditahan", "Antrean", "Terbit",
           "Aturan gagal B0/B1/D", "G1 D", "Dtk"],
          [[d["file"], d["domain"], str(d["halaman"]), "ya" if d["bertag_asli"] else "tidak", str(d["gambar"]), str(d["dekoratif"]),
            str(d["figur_vektor_ditahan"]), str(d["antrean"]), str(d["terbit_otomatis"]),
            f"{d['B0']['aturan_gagal']}/{d['B1']['aturan_gagal']}/{d['D']['aturan_gagal']}",
            "lulus" if d["D"]["g1_lulus"] else "gagal", num(d["detik"], 1)] for d in G6["per_dokumen"]],
          [1300, 1150, 550, 750, 700, 750, 950, 750, 650, 1250, 650, 718], align=["left", "left"] + ["center"] * 10),
    tabel("Tabel N.2. Alamat sumber PDF G6 ringan", ["Dokumen", "URL sumber", "Ukuran (bita)"],
          [[f, cc[f]["url"], f"{cc[f]['bytes']:,}".replace(",", ".")] for f in (d["file"] for d in G6["per_dokumen"]) if f in cc],
          [1500, 7268, 1600]),
)

# O. Rencana data vs realisasi
DJ = json.load(open(PROTO / "data/synth/docs.json", encoding="utf-8"))
CJ = L["charts_json"]
n_rumus = sum(1 for d in DJ if d.get("rumus"))
lampiran(
    "Lampiran O. Rencana Data dan Realisasinya pada MVP",
    "Sumber: Tabel 5 dan Lampiran G (rencana); data/synth/charts.json, data/synth/docs.json, data/cc/cc.json, "
    "eval/results.json (realisasi).",
    tabel("Tabel O.1. Rencana data (Tabel 5) dan realisasi pada evaluasi final", ["Data", "Rencana", "Realisasi", "Dipakai pada"], [
        ["Bingkai publik materi kuliah (Common Crawl, ac.id)", "Sampel berstrata 600 PDF; 60 dokumen G6",
         f"6 PDF dari {G6['sumber']}", "G6 ringan (Lampiran N)"],
        ["Audit awal materi kuliah", "28 PDF, 18 kampus", "28 PDF, 18 kampus (semifinal)", "Bab I"],
        ["Set sintetis penulis", "30 dokumen, 150 grafik, 200 rumus",
         f"{len(DJ)} dokumen ({sum(d['set'] == 'dev' for d in DJ)} dev, {sum(d['set'] == 'uji' for d in DJ)} uji); "
         f"{len(CJ)} grafik ({sum(c['set'] == 'dev' for c in CJ)} dev, {sum(c['set'] == 'uji' for c in CJ)} uji); "
         f"{n_rumus} rumus (satu per dokumen)", "G1, G3, G4 (Lampiran J–M)"],
        ["VisText, ChartQA", "Subset 300 grafik", "Tidak dipakai", "—"],
        ["UniMER-Test, IndoTabVQA", "200 rumus; 300 citra tabel", "Tidak dipakai", "—"],
        ["Korpus SafeDocs [55]", "Pelengkap berkas terpotong", "Tidak dipakai", "—"],
        ["Tolok ukur Kumar dkk.", "Validasi silang G1, G4", "Tidak dipakai", "—"],
    ], [2600, 2300, 3600, 1868]),
    tabel("Tabel O.2. Komposisi jenis grafik pada set sintetis", ["Jenis grafik", "Dev", "Uji", "Uji berlabel nilai"],
          [[JN[j], str(sum(c["set"] == "dev" and c["jenis_grafik"] == j for c in CJ)),
            str(sum(c["set"] == "uji" and c["jenis_grafik"] == j for c in CJ)),
            str(sum(c["set"] == "uji" and c["jenis_grafik"] == j and c["label_nilai"] for c in CJ))] for j in JN],
          [3600, 2200, 2200, 2368], align=["left", "center", "center", "center"]),
)

# P. Antarmuka dan titik akhir
EP = [
    ("GET", "/api/sampel", "Daftar dokumen contoh (set dev, uji, nyata)"),
    ("POST", "/api/sampel/{name}", "Memproses satu dokumen contoh"),
    ("POST", "/api/unggah", "Mengunggah dan memproses PDF"),
    ("GET", "/api/pekerjaan/{jid}", "Status dan hasil satu sesi remediasi"),
    ("DELETE", "/api/pekerjaan/{jid}", "Menghapus sesi"),
    ("GET", "/api/pekerjaan/{jid}/dengar", "Transkrip urutan baca PDF asli dan hasil SIMAK"),
    ("GET", "/api/pekerjaan/{jid}/img/{el}", "Citra elemen untuk antrean tinjauan"),
    ("POST", "/api/tinjau/{jid}/{el}", "Keputusan ULD: setujui, ubah, tolak, dekoratif"),
    ("GET", "/keluaran/{jid}/{name}", "Unduh HTML aksesibel, PDF bertag (D), PDF asli (B0)"),
    ("GET", "/api/demo/grafik", "Daftar grafik uji untuk Demo Gerbang"),
    ("POST", "/api/demo/jalankan/{cid}", "Menjalankan B2 dan SIMAK pada satu grafik"),
    ("POST", "/api/demo/suntik/{cid}", "Menyuntik satu jenis kesalahan dan memeriksa ulang"),
    ("GET", "/api/demo/mutasi", "Daftar jenis injeksi"),
    ("GET", "/api/evaluasi", "Isi eval/results.json untuk halaman Evaluasi"),
    ("GET", "/api/grafik/{cid}.png", "Citra grafik sintetis"),
    ("GET, POST", "/api/suara/daftar, /api/suara, /api/suara/siapkan", "Suara ElevenLabs opsional (aktif bila key diatur)"),
    ("POST", "/api/perintah", "Perintah suara bebas yang dipetakan ke tombol"),
]
lampiran(
    "Lampiran P. Antarmuka MVP dan Titik Akhir Layanan",
    "Sumber: app.py dan web/ pada tag v1.0-final; tangkapan layar diambil luring tanpa API key oleh gambar/tangkap_mvp.py.",
    gambar(GB / "g4_a_remediasi.png", "Gambar P.1. Halaman Remediasi & Tinjauan pada materi_03.pdf (set uji).", 6.0),
    gambar(GB / "g4_b_demo_gerbang.png", "Gambar P.2. Halaman Demo Gerbang pada syn043 setelah injeksi tukar label–nilai.", 6.0),
    gambar(GB / "g4_c_evaluasi.png", "Gambar P.3. Halaman Evaluasi pada set uji tertahan.", 6.0),
    gambar(GB / "g4_d_html_keluaran.png", "Gambar P.4. HTML keluaran dengan bilah mode dengar Web Speech id-ID.", 6.0),
    tabel("Tabel P.1. Titik akhir layanan MVP (app.py)", ["Metode", "Jalur", "Fungsi"], [list(e) for e in EP],
          [1300, 4200, 4868]),
)

sisip(END, [tutup_lanskap, *akhir])

# ------------------------------------------------------------------ simpan dengan pengaman suntingan tangan
if OUT.exists():
    now = hashlib.sha256(OUT.read_bytes()).hexdigest()
    ok = {hashlib.sha256(SRC.read_bytes()).hexdigest(), SIDE.read_text().strip() if SIDE.exists() else ""}
    if now not in ok:
        sys.exit(f"{OUT.name} sudah disunting sejak perakitan terakhir; tidak ditimpa. Hapus atau ganti nama dulu.")
doc.save(OUT)
SIDE.write_text(hashlib.sha256(OUT.read_bytes()).hexdigest())
print("ok", OUT.name)
